"""FileMaid content extractors — local, no OCR / no transcription this pass.

Prefer stdlib + already-in-tree python-docx / openpyxl / python-pptx / Pillow.
PDF uses PyMuPDF (``fitz``) when present; otherwise a install hint.
Never raises. Not registered as @tools.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Union

from primus.core.filemaid_safety import is_path_blocked

PathLike = Union[str, Path]

TEXT_EXTS = frozenset({
    ".txt", ".md", ".markdown", ".json", ".xml", ".csv", ".log",
    ".py", ".js", ".ts", ".html", ".css", ".yaml", ".yml", ".ini", ".cfg",
})
IMAGE_EXTS = frozenset({
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tiff", ".ico",
})
PDF_EXTS = frozenset({".pdf"})
OFFICE_EXTS = frozenset({
    ".docx", ".pptx", ".xlsx", ".xls", ".doc", ".ppt",
})
VIDEO_EXTS = frozenset({
    ".mp4", ".mkv", ".avi", ".mov", ".wmv", ".webm", ".flv", ".m4v",
})
AUDIO_EXTS = frozenset({
    ".mp3", ".wav", ".m4a", ".flac", ".ogg", ".aac", ".wma", ".opus",
})


def detect_file_type(path: PathLike) -> str:
    ext = Path(path).suffix.lower()
    if ext in TEXT_EXTS:
        return "text"
    if ext in PDF_EXTS:
        return "pdf"
    if ext in OFFICE_EXTS:
        return "office"
    if ext in IMAGE_EXTS:
        return "image"
    if ext in VIDEO_EXTS:
        return "video"
    if ext in AUDIO_EXTS:
        return "audio"
    return "other"


@dataclass
class ExtractionOptions:
    enable_ocr: bool = False
    enable_video_keyframes: bool = False
    enable_audio_transcription: bool = False
    max_chars: int = 5000
    max_file_size: int = 500 * 1024 * 1024


def _base_meta(resolved: Path, file_type: str, size_bytes: int, method: str) -> dict[str, Any]:
    return {
        "path": str(resolved),
        "filename": resolved.name,
        "extension": resolved.suffix.lower(),
        "file_type": file_type,
        "size_bytes": size_bytes,
        "extraction_method": method,
        "truncated": False,
    }


def _result(text: str, metadata: dict[str, Any], summary: str = "") -> dict[str, Any]:
    return {"text": text, "metadata": metadata, "summary": summary}


class ContentExtractor:
    def __init__(self, options: ExtractionOptions | None = None) -> None:
        self.options = options or ExtractionOptions()

    def extract(self, file_path: str) -> dict:
        """Return ``{text, metadata, summary}``. Order: blocked → missing → too large → dispatch → truncate."""
        try:
            return self._extract(file_path)
        except Exception as exc:  # noqa: BLE001 — never raise to the caller
            raw = Path(file_path).expanduser()
            meta = _base_meta(raw, detect_file_type(raw), 0, "error")
            meta["error"] = str(exc)
            return _result(f"[Error: {exc}]", meta)

    def _extract(self, file_path: str) -> dict:
        raw = Path(file_path).expanduser()
        blocked, reason = is_path_blocked(raw)
        if blocked:
            try:
                resolved = raw.resolve()
            except Exception:  # noqa: BLE001
                resolved = raw
            meta = _base_meta(resolved, detect_file_type(resolved), 0, "blocked")
            meta["error"] = reason
            return _result(f"[Blocked: {reason}]", meta)

        try:
            resolved = raw.resolve()
        except Exception as exc:  # noqa: BLE001
            meta = _base_meta(raw, detect_file_type(raw), 0, "error")
            meta["error"] = f"Cannot resolve path: {exc}"
            return _result(f"[Error: Cannot resolve path: {exc}]", meta)

        file_type = detect_file_type(resolved)
        if not resolved.exists() or not resolved.is_file():
            meta = _base_meta(resolved, file_type, 0, "missing")
            meta["error"] = f"Not found: {resolved}"
            return _result(f"[Missing: {resolved}]", meta)

        try:
            size_bytes = resolved.stat().st_size
        except Exception as exc:  # noqa: BLE001
            meta = _base_meta(resolved, file_type, 0, "error")
            meta["error"] = str(exc)
            return _result(f"[Error: {exc}]", meta)

        if size_bytes > self.options.max_file_size:
            meta = _base_meta(resolved, file_type, size_bytes, "too_large")
            meta["error"] = (
                f"File too large: {size_bytes} > {self.options.max_file_size}"
            )
            return _result(f"[Too large: {size_bytes} bytes]", meta)

        text, method, extra = self._dispatch(resolved, file_type, size_bytes)
        meta = _base_meta(resolved, file_type, size_bytes, method)
        meta.update(extra)
        if len(text) > self.options.max_chars:
            text = text[: self.options.max_chars]
            meta["truncated"] = True
        return _result(text, meta)

    def _dispatch(
        self, path: Path, file_type: str, size_bytes: int
    ) -> tuple[str, str, dict[str, Any]]:
        if file_type == "text":
            return self._extract_text_file(path)
        if file_type == "pdf":
            return self._extract_pdf(path)
        if file_type == "office":
            return self._extract_office(path)
        if file_type == "image":
            return self._extract_image(path, size_bytes)
        if file_type == "video":
            return (
                f"[Video: {path.suffix.lower()}, {size_bytes} bytes]",
                "av_stat",
                {},
            )
        if file_type == "audio":
            return (
                f"[Audio: {path.suffix.lower()}, {size_bytes} bytes]",
                "av_stat",
                {},
            )
        return f"[Unsupported type: {path.suffix.lower()}]", "other", {}

    def _extract_text_file(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        for enc in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                return path.read_text(encoding=enc), "text", {}
            except (UnicodeDecodeError, UnicodeError):
                continue
            except Exception as exc:  # noqa: BLE001
                return f"[Error: {exc}]", "text", {"error": str(exc)}
        try:
            return path.read_text(encoding="utf-8", errors="replace"), "text", {}
        except Exception as exc:  # noqa: BLE001
            return f"[Error: {exc}]", "text", {"error": str(exc)}

    def _extract_pdf(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        try:
            import fitz  # type: ignore  # PyMuPDF
        except ImportError:
            return "(PDF — install PyMuPDF)", "pdf_missing", {}
        try:
            parts: list[str] = []
            with fitz.open(str(path)) as doc:
                page_count = doc.page_count
                for page in doc:
                    try:
                        parts.append(page.get_text() or "")
                    except Exception:  # noqa: BLE001
                        continue
            text = "\n".join(parts).strip()
            extra: dict[str, Any] = {"page_count": page_count}
            if not text:
                extra["error"] = "No extractable text"
                text = "(PDF has no extractable text)"
            return text, "fitz", extra
        except Exception as exc:  # noqa: BLE001
            return f"[Error: {exc}]", "fitz", {"error": str(exc)}

    def _extract_office(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        ext = path.suffix.lower()
        if ext == ".docx":
            return self._extract_docx(path)
        if ext == ".pptx":
            return self._extract_pptx(path)
        if ext == ".xlsx":
            return self._extract_xlsx(path)
        return (
            f"[Legacy office format {ext} — not extracted]",
            "office_legacy",
            {},
        )

    def _extract_docx(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        try:
            import docx  # python-docx
        except ImportError:
            return "(DOCX — install python-docx)", "docx_missing", {}
        try:
            document = docx.Document(str(path))
            chunks = [p.text for p in document.paragraphs if p.text and p.text.strip()]
            for table in getattr(document, "tables", []):
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                    if cells:
                        chunks.append(" | ".join(cells))
            text = "\n".join(chunks).strip()
            extra: dict[str, Any] = {}
            if not text:
                extra["error"] = "DOCX has no extractable text"
            return text, "docx", extra
        except Exception as exc:  # noqa: BLE001
            return f"[Error: {exc}]", "docx", {"error": str(exc)}

    def _extract_pptx(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        try:
            from pptx import Presentation
        except ImportError:
            return "(PPTX — install python-pptx)", "pptx_missing", {}
        try:
            prs = Presentation(str(path))
            out: list[str] = []
            for idx, slide in enumerate(prs.slides, 1):
                texts: list[str] = []
                for shape in slide.shapes:
                    try:
                        if shape.has_text_frame:
                            for para in shape.text_frame.paragraphs:
                                t = "".join(run.text for run in para.runs).strip()
                                if t:
                                    texts.append(t)
                    except Exception:  # noqa: BLE001
                        continue
                if texts:
                    out.append(f"Slide {idx}:\n" + "\n".join(texts))
            return "\n\n".join(out).strip(), "pptx", {}
        except Exception as exc:  # noqa: BLE001
            return f"[Error: {exc}]", "pptx", {"error": str(exc)}

    def _extract_xlsx(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        try:
            import openpyxl
        except ImportError:
            return "(XLSX — install openpyxl)", "xlsx_missing", {}
        try:
            wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001
            return f"[Error: {exc}]", "xlsx", {"error": str(exc)}
        try:
            rows_out: list[str] = []
            for sn in wb.sheetnames:
                ws = wb[sn]
                rows_out.append(f"## {sn}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 100:
                        rows_out.append("…")
                        break
                    cells = ["" if c is None else str(c) for c in row]
                    if any(cells):
                        rows_out.append(" | ".join(cells))
            return "\n".join(rows_out).strip(), "xlsx", {}
        except Exception as exc:  # noqa: BLE001
            return f"[Error: {exc}]", "xlsx", {"error": str(exc)}
        finally:
            try:
                wb.close()
            except Exception:  # noqa: BLE001
                pass

    def _extract_image(self, path: Path, size_bytes: int) -> tuple[str, str, dict[str, Any]]:
        # OCR is not ported (no easyocr). Flags are accepted and ignored.
        try:
            from PIL import Image
        except ImportError:
            return (
                f"[Image: {path.suffix.lower()}, {size_bytes} bytes]",
                "image_stat",
                {},
            )
        try:
            with Image.open(str(path)) as img:
                width, height = img.size
                mode = img.mode
            return f"[Image: {width}x{height} {mode}]", "pillow", {}
        except Exception as exc:  # noqa: BLE001
            return (
                f"[Image: {path.suffix.lower()}, {size_bytes} bytes]",
                "image_stat",
                {"error": str(exc)},
            )


def extract_text(path: PathLike, max_chars: int = 5000) -> str:
    return ContentExtractor(ExtractionOptions(max_chars=max_chars)).extract(str(path))["text"]
