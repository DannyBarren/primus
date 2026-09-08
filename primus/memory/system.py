"""Primus memory system (Phase 5 extraction).

The whole memory layer lives here: long-term memory (load/save/sync), the multi-collection Chroma
Knowledge Base (ingestion, search, RAG) with its caching embeddings wrapper, document extraction,
the MemorySystem (session + structured LTM + learning digest), reflections/interactions, metrics,
and KB/recall helpers. Moved verbatim from admin_assistant.py; the only mechanical change is that
references to admin_assistant module globals are reached through the live host module as
``_host.<name>`` (keeps the circular import safe and preserves rebound config + singletons).

Singletons (the cached KnowledgeBase / MemorySystem / MetricsTracker, the KB lock and index status)
live here and are reached from elsewhere through their accessor functions, which admin_assistant
re-imports — so call sites are unchanged. Nothing here is locked or read-only.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import threading
import time
import uuid
from collections import OrderedDict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Optional

import sys as _sys

# Live host module (admin_assistant or, when run as the entrypoint, "__main__"); never re-executes.
_host = _sys.modules.get("admin_assistant") or _sys.modules["__main__"]

# Operator identity (generic "the operator" unless ~/.config/primus/operator.json personalizes it).
from primus.core.identity import OPERATOR as _OPERATOR, render as _render_op  # noqa: E402

# langchain surface bound from the host (real objects or graceful stubs resolved at startup).
Document = _host.Document
BaseMessage = _host.BaseMessage
HumanMessage = _host.HumanMessage
SystemMessage = _host.SystemMessage
AIMessage = _host.AIMessage


def load_memory() -> dict[str, Any]:
    _host.ensure_app_dirs()
    try:
        return json.loads(_host.MEMORY_FILE.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {"facts": _host.DEFAULT_BUSINESS_CONTEXT.copy(), "preferences": _host.DEFAULT_PREFERENCES.copy(), "past_tasks": []}


def save_memory(data: dict[str, Any]) -> None:
    data["updated_at"] = _host._now_iso()
    _host.MEMORY_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")


def sync_core_memory() -> None:
    """Merge latest operator facts and default preferences without overwriting user edits."""
    mem = load_memory()
    facts: list[str] = list(mem.get("facts") or [])
    existing = set(facts)
    for fact in _host.DEFAULT_BUSINESS_CONTEXT:
        if fact not in existing:
            facts.append(fact)
    mem["facts"] = facts
    prefs: dict = mem.setdefault("preferences", {})
    for key, value in _host.DEFAULT_PREFERENCES.items():
        prefs.setdefault(key, value)
    save_memory(mem)

# ---------------------------------------------------------------------------
# Knowledge Base (production RAG) — multi-collection Chroma @ ~/.primus/knowledge/
#
# Collections: core_business | technical | projects | personal_preferences | learned
# Each chunk carries metadata: source, kind, collection, tags, importance, indexed_at
# Ingestion: smart chunking (code/markdown/default), hash-based dedup, mtime updates
# ---------------------------------------------------------------------------

_kb_instance: Optional["KnowledgeBase"] = None
_kb_lock = threading.Lock()
_index_status: dict[str, Any] = {"running": False, "last": "", "message": ""}


def _content_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()[:16]


# Lowercased operator project names route seeds/paths into the `projects` collection
# (empty in a stock clone; personalized via ~/.config/primus/operator.json).
_PROJECT_ROUTE_KEYS: tuple[str, ...] = tuple(p.lower() for p in _OPERATOR["projects"])

# Keywords that mark a user message as an identity fact (highest-priority memory). The
# operator's first name is added when the identity is personalized.
_IDENTITY_FACT_KEYS: tuple[str, ...] = (
    "my name", "spelled", "spelling", "creator", "loyalty",
) + (
    (str(_OPERATOR["name"]).split()[0].lower(),) if _OPERATOR["custom"] else ()
)


def classify_kb_collection(
    *,
    source: str,
    kind: str = "manual",
    path: Optional[Path] = None,
    tags: Optional[list[str]] = None,
) -> str:
    """Route documents to the best Chroma collection."""
    src = source.lower()
    if kind == "ltm" or (tags and "preferences" in tags):
        return "personal_preferences"
    if kind in ("web", "manual", "interaction", "tool"):
        return "learned"
    if kind == "seed" or "primus:seed" in src:
        if any(x in src for x in _PROJECT_ROUTE_KEYS):
            return "projects"
        if any(x in src for x in ("toolstack", "hardware", "specialt")):
            return "technical"
        return "core_business"
    if path is not None:
        p = str(path).lower()
        if p.endswith(".py") or "/scripts/" in p or "ai-workshop" in p:
            return "technical"
        if any(x in p for x in _PROJECT_ROUTE_KEYS):
            return "projects"
        if "client" in p or "document" in p:
            return "core_business"
    if tags:
        if "project" in tags:
            return "projects"
        if "business" in tags:
            return "core_business"
    return "learned"


_PDF_PAGE_CAP = 800  # safety cap on pages parsed per PDF


def _extract_pdf_text(path: Path) -> tuple[str, Optional[str]]:
    """Extract text from a PDF using multiple backends for maximum reliability.

    Order: pypdf (lenient) → pdfplumber → PyMuPDF (fitz). Real-world PDFs (exported,
    linearized, or slightly malformed) routinely break a single parser, so we try each
    in turn and return the first that yields text. Never raises.
    """
    errors: list[str] = []

    # 1) pypdf — fast, but strict by default; use strict=False to tolerate minor breakage.
    try:
        from pypdf import PdfReader

        try:
            reader = PdfReader(str(path), strict=False)
        except TypeError:  # very old pypdf without the strict kwarg
            reader = PdfReader(str(path))
        parts: list[str] = []
        for pg in reader.pages[:_PDF_PAGE_CAP]:
            try:
                parts.append(pg.extract_text() or "")
            except Exception as exc:  # noqa: BLE001 — one bad page shouldn't kill the doc
                errors.append(f"pypdf page: {str(exc)[:40]}")
        text = "\n\n".join(p for p in parts if p).strip()
        if text:
            return text, None
        errors.append("pypdf: no embedded text")
    except ImportError:
        errors.append("pypdf not installed")
    except Exception as exc:  # noqa: BLE001
        errors.append(f"pypdf: {str(exc)[:50]}")

    # 2) pdfplumber — handles many PDFs pypdf chokes on (broken xref, complex layouts).
    try:
        import pdfplumber

        parts = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages[:_PDF_PAGE_CAP]:
                try:
                    parts.append(page.extract_text() or "")
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"pdfplumber page: {str(exc)[:40]}")
        text = "\n\n".join(p for p in parts if p).strip()
        if text:
            return text, None
        errors.append("pdfplumber: no embedded text")
    except ImportError:
        errors.append("pdfplumber not installed")
    except Exception as exc:  # noqa: BLE001
        errors.append(f"pdfplumber: {str(exc)[:50]}")

    # 3) PyMuPDF (fitz) — optional, very tolerant; last resort.
    try:
        import fitz  # type: ignore  # PyMuPDF

        parts = []
        with fitz.open(str(path)) as doc:
            for page in list(doc)[:_PDF_PAGE_CAP]:
                try:
                    parts.append(page.get_text() or "")
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"fitz page: {str(exc)[:40]}")
        text = "\n\n".join(p for p in parts if p).strip()
        if text:
            return text, None
        errors.append("fitz: no embedded text")
    except ImportError:
        pass  # optional backend
    except Exception as exc:  # noqa: BLE001
        errors.append(f"fitz: {str(exc)[:50]}")

    # Nothing worked. Distinguish "no reader installed" from "scanned/empty PDF".
    have_reader = any("not installed" not in e for e in errors)
    if not have_reader:
        return "", "No PDF reader available. Run: uv pip install pypdf pdfplumber"
    return "", (
        "No extractable text — likely a scanned/image-only PDF (no embedded text layer; "
        "OCR not available). [" + "; ".join(errors[-2:]) + "]"
    )


def _extract_docx_text(path: Path) -> tuple[str, Optional[str]]:
    """Extract paragraphs + table cells from a .docx. Never raises."""
    try:
        import docx  # python-docx

        document = docx.Document(str(path))
        chunks = [p.text for p in document.paragraphs if p.text and p.text.strip()]
        for table in getattr(document, "tables", []):
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    chunks.append(" | ".join(cells))
        text = "\n".join(chunks).strip()
        return (text, None) if text else ("", "DOCX has no extractable text")
    except ImportError:
        return "", "python-docx not installed. Run: uv pip install python-docx"
    except Exception as exc:  # noqa: BLE001
        return "", f"DOCX read failed: {str(exc)[:80]}"


def _extract_xlsx_text(
    path: Path, *, sheet_name: Optional[str] = None, max_rows: int = 100
) -> tuple[str, Optional[str]]:
    """Extract a readable grid from an Excel .xlsx (openpyxl). Never raises."""
    try:
        import openpyxl
    except ImportError:
        return "", "openpyxl not installed. Run: uv pip install openpyxl"
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        return "", f"XLSX open failed: {str(exc)[:80]}"
    try:
        names = wb.sheetnames
        if sheet_name and sheet_name not in names:
            return "", f"Sheet '{sheet_name}' not found. Available: {', '.join(names)}"
        targets = [sheet_name] if sheet_name else names
        out: list[str] = [f"Workbook `{path.name}` — sheets: {', '.join(names)}"]
        for sn in targets:
            ws = wb[sn]
            out.append(f"\n## Sheet: {sn}  ({ws.max_row}×{ws.max_column})")
            rows: list[str] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows:
                    rows.append(f"… (+{max(0, ws.max_row - max_rows)} more rows)")
                    break
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    rows.append(" | ".join(cells))
            out.append("\n".join(rows) if rows else "(empty)")
        text = "\n".join(out).strip()
        return (text, None) if text else ("", "XLSX has no readable cells")
    except Exception as exc:  # noqa: BLE001
        return "", f"XLSX read failed: {str(exc)[:80]}"
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass


def _extract_pptx_text(path: Path) -> tuple[str, Optional[str]]:
    """Extract slide text + tables from a PowerPoint .pptx (python-pptx). Never raises."""
    try:
        from pptx import Presentation
    except ImportError:
        return "", "python-pptx not installed. Run: uv pip install python-pptx"
    try:
        prs = Presentation(str(path))
        out: list[str] = [f"Presentation `{path.name}` — {len(prs.slides)} slide(s)"]
        for idx, slide in enumerate(prs.slides, 1):
            texts: list[str] = []
            for shape in slide.shapes:
                try:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = "".join(run.text for run in para.runs).strip()
                            if t:
                                texts.append(t)
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            cells = [c.text.strip() for c in row.cells]
                            if any(cells):
                                texts.append(" | ".join(cells))
                except Exception:  # noqa: BLE001 — one odd shape shouldn't drop the slide
                    continue
            out.append(f"\n--- Slide {idx} ---\n" + ("\n".join(texts) if texts else "(no text)"))
        text = "\n".join(out).strip()
        return (text, None) if text else ("", "PPTX has no readable text")
    except Exception as exc:  # noqa: BLE001
        return "", f"PPTX read failed: {str(exc)[:80]}"


def _extract_odt_text(path: Path) -> tuple[str, Optional[str]]:
    """Extract text from an OpenDocument (.odt) via its zip/content.xml — no extra deps."""
    import zipfile

    try:
        with zipfile.ZipFile(str(path)) as zf:
            xml = zf.read("content.xml").decode("utf-8", "replace")
        # Insert spaces at tag boundaries so words don't run together, then strip tags.
        text = re.sub(r"<[^>]+>", " ", xml)
        text = re.sub(r"\s+", " ", text).strip()
        return (text, None) if text else ("", "ODT has no extractable text")
    except Exception as exc:  # noqa: BLE001
        return "", f"ODT read failed: {str(exc)[:80]}"


def _read_text_file(path: Path) -> tuple[str, Optional[str]]:
    """Read a text-like file, trying several encodings before a lossy fallback."""
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=enc), None
        except (UnicodeDecodeError, UnicodeError):
            continue
        except Exception as exc:  # noqa: BLE001
            return "", f"Read failed: {str(exc)[:80]}"
    try:
        return path.read_text(encoding="utf-8", errors="replace"), None
    except Exception as exc:  # noqa: BLE001
        return "", f"Read failed: {str(exc)[:80]}"


def load_document_text(path: Path) -> tuple[str, Optional[str]]:
    """Load plain text from a document. Returns (text, error); never raises.

    Supports PDF (multi-backend), DOCX, ODT, and any text-based format. Unknown or
    extensionless files are attempted as UTF-8 text so nothing is silently dropped.
    """
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return _extract_pdf_text(path)
        if ext == ".docx":
            return _extract_docx_text(path)
        if ext == ".xlsx":
            return _extract_xlsx_text(path)
        if ext == ".pptx":
            return _extract_pptx_text(path)
        if ext == ".odt":
            return _extract_odt_text(path)
        if ext in {".doc", ".rtf"}:
            # Best-effort text scrape; legacy binary/RTF formats need conversion for full fidelity.
            text, err = _read_text_file(path)
            cleaned = re.sub(r"\\[a-z]+\d* ?|[{}]", " ", text) if ext == ".rtf" else text
            cleaned = re.sub(r"[^\x09\x0a\x0d\x20-\x7e]+", " ", cleaned)
            cleaned = re.sub(r"\s+", " ", cleaned).strip()
            if len(cleaned) >= 20:
                return cleaned, None
            return "", (
                f"`{ext}` needs conversion. Save as PDF/DOCX/TXT, or run: "
                f"libreoffice --headless --convert-to pdf '{path.name}'"
            )
        return _read_text_file(path)
    except Exception as exc:  # noqa: BLE001
        return "", str(exc)[:120]


_KW_STOPWORDS = frozenset(
    "the a an of to in on for and or is are was were be been do does did how what "
    "why when where who which my your me i you it this that with about from as at by "
    "can could should would will whats hows tell show give get".split()
)


def _keyword_tokens(text: str) -> set[str]:
    """Lowercase content tokens (alnum, ≥3 chars, no stopwords) for keyword scoring."""
    toks = re.findall(r"[a-z0-9]{3,}", (text or "").lower())
    return {t for t in toks if t not in _KW_STOPWORDS}


def _keyword_overlap_score(query: str, text: str) -> float:
    """BM25-lite: fraction of query content terms present in the candidate text (0..1)."""
    q = _keyword_tokens(query)
    if not q:
        return 0.0
    d = _keyword_tokens(text)
    if not d:
        return 0.0
    return len(q & d) / len(q)


# Business/engineering cues → search core_business + projects (+ technical) first.
# Generic term set; the operator's project names + business tag keywords (operator.json)
# are appended so personalized installs keep their intent-scoped routing.
_BIZ_PROJECT_TERMS: tuple[str, ...] = (
    "consulting", "client", "business", "revenue", "pricing", "invoice", "proposal",
    "offer", "launch", "product", "roadmap", "strateg", "automation", "pipeline",
    "crewai", "wordpress", "ecommerce", "e-commerce", "deploy", "integration", "project",
) + tuple(p.lower() for p in _OPERATOR["projects"]) + tuple(
    w.lower() for w in _OPERATOR["tag_keywords_business"]
)
_BIZ_PROJECT_RE = re.compile(
    r"\b(" + "|".join(re.escape(t) for t in _BIZ_PROJECT_TERMS) + r")\b",
    re.I,
)


def _prefer_collections_for(query: str) -> Optional[list[str]]:
    """Pick high-signal collections to search first for business/engineering queries."""
    if not bool(_host.CFG.get("rag_scope_by_intent", True)):
        return None
    if _BIZ_PROJECT_RE.search(query or ""):
        return ["core_business", "projects", "technical"]
    return None


class _CachingEmbeddings:
    """Wraps a LangChain embeddings object with a bounded in-memory query cache.

    Chroma re-embeds the query text on every similarity search; repeated/identical
    queries (very common in chat) would otherwise re-run the embedding model each time,
    burning CPU and heating the laptop. We memoize embed_query; embed_documents is
    passed straight through (ingestion is one-off and shouldn't be cached).
    """

    def __init__(self, inner: Any, max_size: int = 256) -> None:
        self._inner = inner
        self._cache: "OrderedDict[str, list[float]]" = OrderedDict()
        self._max = max(16, int(max_size))

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return self._inner.embed_documents(texts)

    def embed_query(self, text: str) -> list[float]:
        key = _content_hash((text or "").strip().lower())
        cached = self._cache.get(key)
        if cached is not None:
            self._cache.move_to_end(key)
            return cached
        vec = self._inner.embed_query(text)
        self._cache[key] = vec
        if len(self._cache) > self._max:
            self._cache.popitem(last=False)
        return vec

    def __getattr__(self, name: str) -> Any:
        # Delegate any other attribute access (e.g. internal LangChain hooks) to the wrapped object.
        return getattr(self._inner, name)


class KnowledgeBase:
    """Multi-collection Chroma knowledge store with metadata-rich ingestion."""

    def __init__(self) -> None:
        self._embeddings = None
        self._stores: dict[str, Any] = {}
        self._splitters: dict[str, Any] = {}
        self._embed_failed: bool = False
        self._chroma_cls: Any = None

    @staticmethod
    def _resolve_embedding_choice() -> tuple[str, str]:
        """Return (provider, model) from config. provider ∈ 'hf' | 'ollama'.

        - "auto": prefer light local bge-small (sentence-transformers) → less CPU/heat;
          fall back to Ollama nomic if HF isn't installed.
        - "bge-small"/"bge"/"light": HF BAAI/bge-small-en-v1.5.
        - anything with a "/" (e.g. BAAI/...): HF model id.
        - otherwise: Ollama model name (e.g. nomic-embed-text).
        """
        raw = (_host.CFG.get("embedding_model", "nomic-embed-text") or "").strip()
        light = _host.CFG.get("embedding_fallback", "BAAI/bge-small-en-v1.5")
        low = raw.lower()
        if low in ("auto", ""):
            return ("hf", light) if _host._module_available("sentence_transformers") else ("ollama", "nomic-embed-text")
        if low in ("bge-small", "bge", "light", "bge_small"):
            return ("hf", light)
        if "/" in raw:  # explicit HF repo id, e.g. BAAI/bge-small-en-v1.5
            return ("hf", raw)
        # Bare HF-style sentence-transformer names (no slash) → HF. Normalize the common
        # bge-small alias to its canonical repo id so it resolves correctly.
        if any(tok in low for tok in ("bge", "minilm", "all-mini", "e5-", "gte-", "mpnet", "sentence-transformers")):
            if low.startswith("bge-small") or low == "bge-small-en-v1.5":
                return ("hf", "BAAI/bge-small-en-v1.5")
            return ("hf", raw)
        return ("ollama", raw)

    def _get_embeddings(self):
        if self._embeddings is not None:
            return self._embeddings
        if self._embed_failed:
            raise RuntimeError(
                "Knowledge base embeddings unavailable. Run: ollama pull nomic-embed-text "
                "OR uv pip install sentence-transformers"
            )
        provider, model = self._resolve_embedding_choice()
        cache_size = int(_host.CFG.get("embedding_cache_size", 256))

        def _wrap(emb: Any) -> Any:
            return _CachingEmbeddings(emb, max_size=cache_size) if cache_size > 0 else emb

        # Preferred provider first, then the other as fallback — order depends on config.
        order = ["hf", "ollama"] if provider == "hf" else ["ollama", "hf"]
        last_exc: Optional[Exception] = None
        for prov in order:
            if prov == "ollama":
                try:
                    from langchain_ollama import OllamaEmbeddings

                    _host._ensure_ollama_host_env()
                    name = model if provider == "ollama" else "nomic-embed-text"
                    emb = OllamaEmbeddings(model=name)
                    emb.embed_query("primus-kb-init")
                    self._embeddings = _wrap(emb)
                    _host.log.info("KB embeddings: Ollama %s (cache=%d)", name, cache_size)
                    return self._embeddings
                except Exception as exc:  # noqa: BLE001
                    last_exc = exc
                    _host.log.warning("Ollama embeddings unavailable (%s)", exc)
            else:  # hf
                for hf_model in (
                    model if provider == "hf" else _host.CFG.get("embedding_fallback", "BAAI/bge-small-en-v1.5"),
                    "sentence-transformers/all-MiniLM-L6-v2",
                ):
                    try:
                        try:
                            from langchain_huggingface import HuggingFaceEmbeddings  # type: ignore
                        except ImportError:
                            from langchain_community.embeddings import HuggingFaceEmbeddings
                        self._embeddings = _wrap(HuggingFaceEmbeddings(model_name=hf_model))
                        _host.log.info("KB embeddings: HuggingFace %s (cache=%d)", hf_model, cache_size)
                        return self._embeddings
                    except Exception as exc:  # noqa: BLE001
                        last_exc = exc
                        _host.log.warning("HF embed %s failed: %s", hf_model, exc)
        self._embed_failed = True
        raise RuntimeError(
            f"Knowledge base needs embeddings ({last_exc}). Run: ollama pull nomic-embed-text "
            "OR uv pip install sentence-transformers"
        )

    def _get_chroma_class(self):
        if self._chroma_cls is not None:
            return self._chroma_cls
        try:
            from langchain_chroma import Chroma  # type: ignore

            self._chroma_cls = Chroma
            return self._chroma_cls
        except ImportError as exc:
            _host.log.error(
                "langchain-chroma required for RAG: uv pip install langchain-chroma (%s)",
                exc,
            )
            raise RuntimeError(
                "Install langchain-chroma: uv pip install langchain-chroma chromadb"
            ) from exc

    def _get_collection(self, name: str):
        if name not in _host.KB_COLLECTIONS:
            name = "learned"
        if name not in self._stores:
            Chroma = self._get_chroma_class()
            _host.KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)
            _host.CHROMA_DIR.mkdir(parents=True, exist_ok=True)
            self._stores[name] = Chroma(
                collection_name=f"primus_{name}",
                persist_directory=str(_host.CHROMA_DIR),
                embedding_function=self._get_embeddings(),
            )
        return self._stores[name]

    def _get_store(self):
        """Backward compat — default collection."""
        return self._get_collection("learned")

    def warm_collections(self) -> None:
        for name in _host.KB_COLLECTIONS:
            self._get_collection(name)

    def _splitter_for(self, path: Optional[Path] = None):
        key = path.suffix.lower() if path else "default"
        if key in self._splitters:
            return self._splitters[key]
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        size = int(_host.CFG.get("rag_chunk_size", 900))
        overlap = int(_host.CFG.get("rag_chunk_overlap", 120))
        if key == ".py":
            splitter = RecursiveCharacterTextSplitter.from_language(
                "python", chunk_size=1000, chunk_overlap=80
            )
        elif key in {".md", ".markdown"}:
            splitter = RecursiveCharacterTextSplitter.from_language(
                "markdown", chunk_size=size, chunk_overlap=overlap
            )
        else:
            splitter = RecursiveCharacterTextSplitter(
                chunk_size=size,
                chunk_overlap=overlap,
                separators=["\n\n", "\n", ". ", " ", ""],
            )
        self._splitters[key] = splitter
        return splitter

    def _load_manifest(self) -> dict[str, Any]:
        _host.ensure_app_dirs()
        try:
            data = json.loads(_host.KB_MANIFEST.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
        except (json.JSONDecodeError, OSError):
            return {}

    def _save_manifest(self, data: dict[str, Any]) -> None:
        data["last_updated"] = _host._now_iso()
        _host.KB_MANIFEST.write_text(json.dumps(data, indent=2), encoding="utf-8")

    def _chunk_documents(
        self,
        text: str,
        *,
        source: str,
        kind: str,
        collection: str,
        path: Optional[Path] = None,
        importance: float = 0.5,
        tags: Optional[list[str]] = None,
        extra: Optional[dict] = None,
    ) -> tuple[list[Document], list[str]]:
        text = text.strip()
        if not text:
            return [], []
        tags = tags or detect_memory_tags(text)
        meta_base = {
            "source": source,
            "kind": kind,
            "collection": collection,
            "tags": ",".join(tags[:6]),
            "importance": float(importance),
            "indexed_at": _host._now_iso(),
            "content_hash": _content_hash(text),
        }
        if extra:
            meta_base.update(extra)
        # Split anything larger than one chunk so chunk_size is actually respected (keeps
        # retrieval precise and embeddings cheap). Tiny texts stay whole.
        split_threshold = int(_host.CFG.get("rag_chunk_size", 600)) + int(_host.CFG.get("rag_chunk_overlap", 80))
        if len(text) <= split_threshold:
            chunks = [text]
        else:
            chunks = self._splitter_for(path).split_text(text)
        docs: list[Document] = []
        ids: list[str] = []
        for i, chunk in enumerate(chunks):
            doc_id = str(uuid.uuid4())
            meta = {**meta_base, "doc_id": doc_id, "chunk": i}
            docs.append(Document(page_content=chunk, metadata=meta))
            ids.append(doc_id)
        return docs, ids

    def _delete_by_source(self, source: str, collection: Optional[str] = None) -> int:
        cols = [collection] if collection else list(_host.KB_COLLECTIONS)
        removed = 0
        for col in cols:
            try:
                vs = self._get_collection(col)
                existing = vs._collection.get(where={"source": source})
                ids = existing.get("ids") or []
                if ids:
                    vs._collection.delete(ids=ids)
                    removed += len(ids)
            except Exception as exc:
                _host.log.debug("Delete source %s in %s: %s", source, col, exc)
        return removed

    def learn_text(
        self,
        text: str,
        *,
        source: str = "user:manual",
        kind: str = "manual",
        collection: Optional[str] = None,
        importance: float = 0.7,
        tags: Optional[list[str]] = None,
        replace_source: bool = False,
        path: Optional[Path] = None,
    ) -> str:
        collection = collection or classify_kb_collection(
            source=source, kind=kind, path=path, tags=tags
        )
        with _kb_lock:
            if replace_source:
                self._delete_by_source(source, collection)
            docs, ids = self._chunk_documents(
                text,
                source=source,
                kind=kind,
                collection=collection,
                path=path,
                importance=importance,
                tags=tags,
            )
            if not docs:
                return "Nothing to store (empty text)."
            self._get_collection(collection).add_documents(docs, ids=ids)
            manifest = self._load_manifest()
            manifest.setdefault("sources", {})[source] = {
                "collection": collection,
                "chunks": len(docs),
                "importance": importance,
                "indexed_at": _host._now_iso(),
            }
            self._save_manifest(manifest)
            return f"Stored {len(docs)} chunk(s) → `{collection}` from `{source}`"

    def embeddings_ready(self) -> tuple[bool, str]:
        """Preflight: is the embedding backend usable? (Avoids per-file cryptic failures.)"""
        try:
            self._get_embeddings()
            return True, ""
        except Exception as exc:  # noqa: BLE001
            return False, str(exc)[:200]

    def _ingest_single(
        self,
        path: Path,
        *,
        incremental: bool = True,
        collection: Optional[str] = None,
        tags: Optional[list[str]] = None,
        extra_meta: Optional[dict] = None,
    ) -> dict[str, Any]:
        """Ingest one file and return a STRUCTURED result (never raises).

        Returns: {status, name, chunks, collection, removed, error}
          status ∈ indexed | already | empty | unreadable | unsupported | missing | error
        """
        res: dict[str, Any] = {
            "status": "error", "name": getattr(path, "name", str(path)),
            "chunks": 0, "collection": collection or "", "removed": 0, "error": "",
        }
        try:
            path = path.expanduser().resolve()
            res["name"] = path.name
            if not path.exists():
                res["status"] = "missing"
                res["error"] = "file not found"
                return res
            if path.is_dir():
                res["status"] = "unsupported"
                res["error"] = "is a directory"
                return res
            if path.suffix.lower() not in _host.TEXT_INGEST_EXTENSIONS:
                res["status"] = "unsupported"
                res["error"] = f"unsupported type {path.suffix or '(none)'}"
                return res
            from primus.core.vault import is_ingest_blocked  # noqa: PLC0415

            if is_ingest_blocked(path):
                res["status"] = "unsupported"
                res["error"] = "oauth/vault path — not ingested"
                return res

            source = f"file:{path}"
            mtime = path.stat().st_mtime
            raw, err = load_document_text(path)
            if err:
                res["status"] = "unreadable"
                res["error"] = err
                return res
            if not raw or not raw.strip():
                res["status"] = "empty"
                res["error"] = "no extractable text"
                return res

            content_hash = _content_hash(raw)
            manifest = self._load_manifest()
            files = manifest.setdefault("files", {})
            prev = files.get(source)
            if incremental and prev and prev.get("mtime") == mtime and prev.get("content_hash") == content_hash:
                res["status"] = "already"
                res["chunks"] = int(prev.get("chunks", 0))
                res["collection"] = prev.get("collection", "")
                return res

            col = collection or classify_kb_collection(source=source, kind="file", path=path)
            res["collection"] = col
            extra = {"path": str(path), "filename": path.name}
            if extra_meta:
                extra.update({k: v for k, v in extra_meta.items() if v not in (None, "")})
            with _kb_lock:
                removed = self._delete_by_source(source, col)
                docs, ids = self._chunk_documents(
                    raw, source=source, kind="file", collection=col, path=path,
                    importance=0.6, tags=tags, extra=extra,
                )
                if not docs:
                    res["status"] = "empty"
                    res["error"] = "empty after chunking"
                    return res
                self._get_collection(col).add_documents(docs, ids=ids)
                files[source] = {
                    "mtime": mtime, "content_hash": content_hash, "collection": col,
                    "chunks": len(docs), "path": str(path),
                }
                self._save_manifest(manifest)
            res["status"] = "indexed"
            res["chunks"] = len(docs)
            res["removed"] = removed
            return res
        except Exception as exc:  # noqa: BLE001 — one bad file must never crash a batch
            _host.log.exception("Ingest failed for %s", res["name"])
            res["status"] = "error"
            res["error"] = str(exc)[:160]
            return res

    def ingest_file(
        self,
        path: Path,
        *,
        incremental: bool = True,
        collection: Optional[str] = None,
        tags: Optional[list[str]] = None,
        extra_meta: Optional[dict] = None,
    ) -> str:
        path = path.expanduser().resolve()
        if path.is_dir():
            return self.ingest_folder(path, incremental=incremental, collection=collection)
        r = self._ingest_single(
            path, incremental=incremental, collection=collection,
            tags=tags, extra_meta=extra_meta,
        )
        status = r["status"]
        if status == "indexed":
            msg = f"Indexed `{r['name']}` → {r['chunks']} chunk(s) in `{r['collection']}`"
            if r["removed"]:
                msg += f" (updated; removed {r['removed']} old chunk(s))"
            return msg
        if status == "already":
            return f"Already indexed (unchanged): `{r['name']}`"
        if status == "missing":
            return f"Not found: {path}"
        if status == "unsupported":
            return f"Skipped unsupported type: {path.suffix or '(none)'}"
        if status == "empty":
            return f"Empty file: `{r['name']}`"
        if status == "unreadable":
            return f"Could not read `{r['name']}`: {r['error']}"
        return f"Failed `{r['name']}`: {r['error']}"

    def ingest_folder(
        self,
        folder: Path,
        *,
        incremental: bool = True,
        collection: Optional[str] = None,
    ) -> str:
        folder = folder.expanduser().resolve()
        if not folder.is_dir():
            return f"Not a directory: {folder}"
        indexed = skipped = errors = 0
        for root, dirnames, filenames in os.walk(folder):
            dirnames[:] = [d for d in dirnames if d not in _host.SKIP_DIR_NAMES]
            for name in filenames:
                fp = Path(root) / name
                if fp.suffix.lower() not in _host.TEXT_INGEST_EXTENSIONS:
                    skipped += 1
                    continue
                try:
                    result = self.ingest_file(fp, incremental=incremental, collection=collection)
                    if result.startswith("Indexed"):
                        indexed += 1
                    elif result.startswith("Already"):
                        skipped += 1
                except Exception as exc:
                    errors += 1
                    _host.log.warning("Ingest failed for %s: %s", fp, exc)
        return f"Folder `{folder}`: indexed {indexed}, skipped {skipped}, errors {errors}."

    def search(
        self,
        query: str,
        k: Optional[int] = None,
        *,
        collection: Optional[str] = None,
        kind: Optional[str] = None,
        exclude_kinds: Optional[set[str]] = None,
        prefer_collections: Optional[list[str]] = None,
    ) -> list[Document]:
        """Hybrid (vector + BM25-style keyword) search, optionally scoped by collection.

        - `collection`: restrict to a single collection.
        - `prefer_collections`: search these first (e.g. core_business/projects for
          business queries); broaden to the rest only if we don't have enough hits.
          Reduces CPU and topic drift by keeping retrieval on-domain.
        """
        k = k or int(_host.CFG.get("kb_search_k", _host.CFG.get("rag_top_k", 4)))
        q = query.strip()
        if not q:
            return []

        if collection:
            ordered_cols = [collection]
        elif prefer_collections and bool(_host.CFG.get("rag_scope_by_intent", True)):
            rest = [c for c in _host.KB_COLLECTIONS if c not in prefer_collections]
            ordered_cols = [c for c in prefer_collections if c in _host.KB_COLLECTIONS] + rest
        else:
            ordered_cols = list(_host.KB_COLLECTIONS)

        # Pull a few extra candidates so the keyword reranker has room to reorder.
        cand_k = min(15, max(k, k * 3)) if bool(_host.CFG.get("hybrid_retrieval", True)) else k

        def _gather(cols: list[str]) -> list[tuple[float, Document]]:
            out: list[tuple[float, Document]] = []
            for col in cols:
                try:
                    vs = self._get_collection(col)
                    for doc, score in vs.similarity_search_with_score(q, k=cand_k):
                        doc.metadata["collection"] = col
                        out.append((float(score), doc))
                except Exception as exc:  # noqa: BLE001
                    _host.log.debug("KB search in %s failed: %s", col, exc)
            return out

        # Scoped first; broaden only if the preferred scope underfills.
        if prefer_collections and bool(_host.CFG.get("rag_scope_by_intent", True)) and not collection:
            scoped = [c for c in prefer_collections if c in _host.KB_COLLECTIONS]
            scored = _gather(scoped)
            if len({id(d) for _, d in scored}) < k:
                scored += _gather([c for c in ordered_cols if c not in scoped])
        else:
            scored = _gather(ordered_cols)

        def _filt(docs: list[Document]) -> list[Document]:
            if kind:
                docs = [d for d in docs if d.metadata.get("kind") == kind]
            if exclude_kinds:
                docs = [d for d in docs if d.metadata.get("kind") not in exclude_kinds]
            return docs

        if not bool(_host.CFG.get("hybrid_retrieval", True)):
            scored.sort(key=lambda x: x[0])
            return _filt([d for _, d in scored])[:k]

        # --- Hybrid rerank: reciprocal-rank fusion of vector distance + keyword overlap ---
        vec_sorted = sorted(scored, key=lambda x: x[0])
        ref: dict[int, Document] = {}
        rrf: dict[int, float] = {}
        for rank, (_, doc) in enumerate(vec_sorted):
            ref[id(doc)] = doc
            rrf[id(doc)] = rrf.get(id(doc), 0.0) + 1.0 / (60 + rank)
        kw_sorted = sorted(
            scored, key=lambda x: _keyword_overlap_score(q, x[1].page_content), reverse=True
        )
        for rank, (_, doc) in enumerate(kw_sorted):
            if _keyword_overlap_score(q, doc.page_content) > 0:  # only reward real overlap
                rrf[id(doc)] = rrf.get(id(doc), 0.0) + 0.5 / (60 + rank)
        fused = sorted(ref.values(), key=lambda d: rrf.get(id(d), 0.0), reverse=True)
        return _filt(fused)[:k]

    def format_results(self, docs: list[Document], *, title: str = "Knowledge base results") -> str:
        if not docs:
            return "*(No matching knowledge base entries.)*"
        lines = [f"**{title}**", ""]
        for i, doc in enumerate(docs, 1):
            meta = doc.metadata or {}
            src = meta.get("source", "unknown")
            kind = meta.get("kind", "text")
            col = meta.get("collection", "?")
            imp = meta.get("importance", "")
            imp_s = f" imp={imp}" if imp else ""
            lines.append(
                f"**[KB-{i}]** `{col}` · `{src}` ({kind}{imp_s})\n{doc.page_content.strip()[:700]}"
            )
        return "\n\n".join(lines)

    def format_sources_block(self, docs: list[Document]) -> str:
        if not docs:
            return ""
        lines = ["**Sources from Knowledge Base:**"]
        for i, doc in enumerate(docs, 1):
            meta = doc.metadata or {}
            lines.append(
                f"- [KB-{i}] {meta.get('collection', '?')} — `{meta.get('source', '?')}`"
            )
        return "\n".join(lines)

    def forget(self, query: str, *, protect_kinds: Optional[set[str]] = None) -> str:
        protect_kinds = protect_kinds or {"seed"}
        docs = self.search(query, k=20)
        ids_by_col: dict[str, list[str]] = {}
        for doc in docs:
            if doc.metadata.get("kind") in protect_kinds:
                continue
            doc_id = doc.metadata.get("doc_id")
            col = doc.metadata.get("collection", "learned")
            if doc_id:
                ids_by_col.setdefault(col, []).append(doc_id)
        total = 0
        with _kb_lock:
            for col, ids in ids_by_col.items():
                self._get_collection(col).delete(ids=ids)
                total += len(ids)
        if not total:
            return "No matching entries to remove (seed docs protected)."
        return f"Removed {total} chunk(s) matching your query."

    def stats(self) -> dict[str, Any]:
        manifest = self._load_manifest()
        files = manifest.get("files") or {}
        collections: dict[str, int] = {}
        total = 0
        for name in _host.KB_COLLECTIONS:
            try:
                count = self._get_collection(name)._collection.count()
            except Exception:
                count = 0
            collections[name] = count
            total += count
        disk_mb = 0.0
        try:
            disk_mb = sum(
                f.stat().st_size for f in _host.CHROMA_DIR.rglob("*") if f.is_file()
            ) / (1024 * 1024)
        except OSError:
            pass
        mem = load_memory()
        return {
            "total_chunks": total,
            "collections": collections,
            "chunks": total,
            "indexed_files": len(files),
            "facts_json": len(mem.get("facts") or []),
            "preferences_json": len(mem.get("preferences") or {}),
            "interactions_json": len(_load_interactions()),
            "chroma_path": str(_host.CHROMA_DIR),
            "embedding_model": _host.CFG.get("embedding_model", "nomic-embed-text"),
            "disk_mb": round(disk_mb, 2),
            "last_ingest": manifest.get("last_updated"),
        }

    def sync_seed_knowledge(self) -> None:
        manifest = self._load_manifest()
        if manifest.get("seed_version", 0) >= _host.SEED_KNOWLEDGE_VERSION:
            return
        for item in _host.SEED_KNOWLEDGE:
            col = classify_kb_collection(source=item["source"], kind="seed")
            self.learn_text(
                item["text"],
                source=item["source"],
                kind="seed",
                collection=col,
                importance=0.95,
                tags=["business", "project"],
                replace_source=True,
            )
        manifest["seed_version"] = _host.SEED_KNOWLEDGE_VERSION
        self._save_manifest(manifest)
        _host.log.info("Synced seed knowledge v%s into collections", _host.SEED_KNOWLEDGE_VERSION)

    def seed_if_empty(self) -> None:
        try:
            if self.stats()["total_chunks"] >= len(_host.SEED_KNOWLEDGE):
                self.sync_seed_knowledge()
                return
        except Exception:
            pass
        self.sync_seed_knowledge()


def get_kb() -> KnowledgeBase:
    global _kb_instance
    if _kb_instance is None:
        _kb_instance = KnowledgeBase()
    return _kb_instance


def render_kb_dashboard() -> str:
    """Markdown dashboard for Menu → Knowledge and /kb status."""
    try:
        kb = get_kb()
        st = kb.stats()
    except Exception as exc:
        return f"**Knowledge Base unavailable:** {exc}"
    idx = _index_status
    lines = [
        "## Knowledge Base",
        "",
        f"- **Embeddings:** `{st['embedding_model']}`",
        f"- **Storage:** `{st['chroma_path']}` ({st.get('disk_mb', 0)} MB)",
        f"- **Total chunks:** {st['total_chunks']} · **Files indexed:** {st['indexed_files']}",
        f"- **Last manifest update:** {st.get('last_ingest') or '—'}",
    ]
    if idx.get("running"):
        lines.append(f"- **Indexing:** in progress — {idx.get('message', '')}")
    elif idx.get("last"):
        lines.append(f"- **Last bulk index:** {idx.get('last')}")
    lines.extend(["", "### Collections", "", "| Collection | Chunks |", "|------------|--------|"])
    for name in _host.KB_COLLECTIONS:
        lines.append(f"| `{name}` | {st['collections'].get(name, 0)} |")
    pending = len(_host.PrimusSession.pending_kb_learn)
    if pending:
        lines.append(f"\n> **{pending}** web result(s) queued — use `/approve learn` to save.")
    lines.extend([
        "",
        "### Commands",
        "`/kb status` · `/kb search query` · `/kb learn text` · `/index_projects` · `/learn ~/path`",
    ])
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Drag-and-drop Knowledge Base ingestion
#
# UI-facing category labels map to the underlying Chroma collections. Projects get an
# extra sub-tag (from the operator's project list) stored as metadata so they can be filtered later.
# Uploaded originals are copied into ~/.primus/knowledge/uploads/<collection>[/<project>]
# so the source path is stable and the document can be re-ingested or audited later.
# ---------------------------------------------------------------------------

KB_CATEGORY_CHOICES: tuple[tuple[str, str], ...] = (
    ("Core Business", "core_business"),
    ("Technical / Linux", "technical"),
    ("Projects", "projects"),
    ("Personal Preferences", "personal_preferences"),
    ("Learned / New Information", "learned"),
)
KB_CATEGORY_LABELS: tuple[str, ...] = tuple(label for label, _ in KB_CATEGORY_CHOICES)
KB_CATEGORY_TO_COLLECTION: dict[str, str] = {label: col for label, col in KB_CATEGORY_CHOICES}
KB_PROJECT_CHOICES: tuple[str, ...] = tuple(_OPERATOR["projects"]) + ("Other",)


from primus.utils.text import safe_filename as _safe_filename  # noqa: E402


def ingest_uploaded_files(
    file_paths: list[str],
    category_label: str,
    project: str = "",
    *,
    progress_cb: Optional[_host.Callable[[float, str], None]] = None,
    extra_tags: Optional[list[str]] = None,
) -> str:
    """Ingest drag-and-dropped files into the chosen KB collection with rich metadata.

    Copies each upload to a stable location, extracts text (PDF/docx/txt/md/py/…),
    chunks + embeds via the existing pipeline, and tags it with category + project.
    Returns a clean, human-readable summary.
    """
    paths = [p for p in (file_paths or []) if p]
    if not paths:
        return "No files dropped. Drag documents into the zone, pick a category, then Process All."

    collection = KB_CATEGORY_TO_COLLECTION.get(category_label, "learned")
    is_projects = collection == "projects"
    proj = (project or "").strip()

    try:
        kb = get_kb()
    except Exception as exc:  # noqa: BLE001
        return f"⚠ Knowledge base unavailable: {exc}"

    # Preflight the embedding backend ONCE so we don't fail every file with a cryptic error.
    ok, emb_err = kb.embeddings_ready()
    if not ok:
        return (
            "⚠ Can't ingest — the embedding model isn't available, so nothing can be indexed.\n\n"
            f"`{emb_err}`\n\n"
            "Fix: make sure Ollama is running and run `ollama pull nomic-embed-text` "
            "(or `uv pip install sentence-transformers` for the local model), then try again."
        )

    try:
        dest_dir = _host.KB_UPLOADS_DIR / collection / (_safe_filename(proj) if (is_projects and proj) else "")
        dest_dir.mkdir(parents=True, exist_ok=True)
    except Exception as exc:  # noqa: BLE001
        _host.log.warning("Upload dir create failed: %s", exc)
        dest_dir = None  # fall back to ingesting temp paths directly

    tags = ["upload"]
    if is_projects and proj:
        tags.append(proj.lower())
    # Project-chat tagging (metadata-level): lets a document be associated with a chat scope.
    for t in (extra_tags or []):
        t = str(t).strip().lower()
        if t and t not in tags:
            tags.append(t)
    extra_meta = {
        "category": category_label,
        "project": proj if is_projects else "",
        "uploaded_at": _host._now_iso(),
        "ingest_source": "drag-and-drop",
    }

    total = len(paths)
    added_files = 0
    added_chunks = 0
    skipped: list[str] = []
    failed: list[str] = []
    detail: list[str] = []
    used_names: set[str] = set()

    for i, src in enumerate(paths):
        src_path = Path(str(src)).expanduser()
        fname = _safe_filename(src_path.name)
        if progress_cb:
            progress_cb(i / max(total, 1), f"Processing {fname} ({i + 1}/{total})")

        if not src_path.exists():
            failed.append(f"{fname}: upload not found (re-drop it)")
            continue
        if src_path.suffix.lower() not in _host.TEXT_INGEST_EXTENSIONS:
            skipped.append(fname)
            continue

        # Copy to a stable, collision-free location so re-ingest/audit is possible.
        dest = src_path
        if dest_dir is not None:
            unique = fname
            if unique in used_names or (dest_dir / unique).exists():
                stem, ext = os.path.splitext(fname)
                unique = f"{stem}_{_content_hash(str(src_path))[:6]}{ext}"
            used_names.add(unique)
            target = dest_dir / unique
            try:
                shutil.copy2(src_path, target)
                dest = target
            except Exception as exc:  # noqa: BLE001
                _host.log.warning("Upload copy failed for %s: %s — ingesting temp path", fname, exc)

        r = kb._ingest_single(
            dest, incremental=False, collection=collection, tags=tags, extra_meta=extra_meta,
        )
        status = r["status"]
        if status == "indexed":
            added_files += 1
            added_chunks += r["chunks"]
            detail.append(f"• {fname} → {r['chunks']} chunk(s)")
        elif status == "already":
            added_files += 1
            detail.append(f"• {fname} → already indexed")
        elif status in ("empty", "unreadable"):
            failed.append(f"{fname}: {r['error'][:70]}")
        elif status == "unsupported":
            skipped.append(fname)
        else:
            failed.append(f"{fname}: {r['error'][:70]}")

    if progress_cb:
        progress_cb(1.0, "Done")

    where = category_label + (f" / {proj}" if (is_projects and proj) else "")
    if added_files:
        head = f"✅ Added {added_files} document(s) ({added_chunks} chunks) to **{where}**."
    elif failed or skipped:
        head = f"⚠ Nothing was added to **{where}** — see details below."
    else:
        head = f"No documents processed for **{where}**."
    parts = [head]
    if detail:
        parts.append("\n".join(detail[:25]))
    if skipped:
        parts.append(f"Skipped (unsupported type): {', '.join(skipped[:12])}")
    if failed:
        parts.append(f"⚠ Failed ({len(failed)}): " + " · ".join(failed[:12]))
    return "\n\n".join(parts)


def handle_kb_command(raw: str) -> Optional[str]:
    """Parse /kb subcommands; returns response text or None if not a /kb command."""
    low = raw.strip().lower()
    if not low.startswith("/kb"):
        return None
    parts = raw.strip().split(None, 2)
    sub = parts[1].lower() if len(parts) > 1 else "status"
    arg = parts[2] if len(parts) > 2 else (parts[1] if len(parts) > 1 and parts[1].lower() not in (
        "status", "search", "learn", "stats",
    ) else "")
    if len(parts) == 2 and parts[1].lower() not in ("status", "search", "learn", "stats"):
        sub, arg = "search", parts[1]

    if sub in ("status", "stats"):
        return render_kb_dashboard()
    if sub == "search":
        if not arg:
            return "Usage: `/kb search your query`"
        docs = get_kb().search(arg)
        return get_kb().format_results(docs, title=f"Search: {arg}")
    if sub == "learn":
        if not arg:
            return "Usage: `/kb learn your note` or `/kb learn ~/path/to/doc.md`"
        if arg.startswith("~") or arg.startswith("/") or arg.startswith("."):
            path = _host._resolve_path(arg.split()[0])
            return get_kb().ingest_file(path)
        return get_kb().learn_text(arg, source="user:/kb learn", importance=0.85)
    return render_kb_dashboard()


def index_projects(*, background: bool = True) -> str:
    paths = list(_host.CFG.get("ingest_paths") or [])
    if not paths:
        paths = [str(Path(p).expanduser()) for p in _host.PROJECT_INDEX_PATHS]
    if background:
        start_background_index(paths)
        return (
            "**Project indexing started** (background)\n\n"
            + "\n".join(f"- `{p}`" for p in paths)
        )
    reports = [get_kb().ingest_folder(Path(p), incremental=True) for p in paths if Path(p).exists()]
    return "\n".join(reports) or "No paths found."

def init_knowledge_base(*, seed: bool = True, auto_index: Optional[bool] = None) -> None:
    """Initialize Chroma, seed business context, optionally background-index configured paths."""
    _host.ensure_app_dirs()
    sync_core_memory()
    try:
        kb = get_kb()
        kb.warm_collections()
        if seed:
            kb.seed_if_empty()
    except Exception as exc:
        _host.log.warning("Knowledge base init deferred: %s", exc)
        return

    if auto_index is None:
        auto_index = bool(_host.CFG.get("auto_index_on_start", False))
    if auto_index:
        start_background_index()
    init_memory_system()


# ---------------------------------------------------------------------------
# Anti-hallucination + unified retrieval (Primus & Forge share this layer)
# Injected into {rag_context} before every agent response.
# ---------------------------------------------------------------------------

ANTI_HALLUCINATION_RULES = _render_op("""
## Anti-hallucination protocol (MANDATORY — Primus & Forge)

1. **Retrieval first** — The blocks below are your verified context. Read them before answering.
2. **Prefer retrieval over generation** — If [KB-N], [MEM-N], or [CONV-N] cover the question, use them.
3. **Cite sources in your reply:**
   - **From knowledge base:** [KB-N] …
   - **From long-term memory:** [MEM-N] …
   - **From previous conversations:** [CONV-N] or "As we discussed {date_label} about …"
4. **Natural continuity** — When [CONV-N] or session summary mentions a project, reference it naturally
   (e.g. "As we discussed last week about the {op_project_example} pipeline…") — only if retrieval supports it.
5. **Unknown = say so** — If no retrieval matches and tools return nothing, say:
   "I don't have that in my knowledge base or past conversations — I can search or you can /learn it."
6. **Never invent** — client names, file paths, repo details, or decisions not in retrieval/tools.
7. **Label inference** — Anything not from retrieval/tools must be prefixed **Reasoning:**.
8. **Domain facts** — Use search_knowledge / recall_memory tools if retrieval looks incomplete.
""")


def _temporal_label(iso_date: str) -> str:
    """Human-friendly time reference for natural recall phrasing."""
    try:
        dt = datetime.fromisoformat(iso_date.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        now = datetime.now(timezone.utc)
        delta = now - dt
        days = delta.days
        if days == 0:
            return "earlier today"
        if days == 1:
            return "yesterday"
        if days < 7:
            return f"{days} days ago"
        if days < 14:
            return "last week"
        if days < 45:
            return "last month"
        return dt.strftime("%B %Y")
    except (ValueError, TypeError):
        return "previously"


def _load_conversation_archive() -> list[dict]:
    _host.ensure_app_dirs()
    try:
        data = json.loads(_host.CONVERSATION_ARCHIVE.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError):
        return []


def _save_conversation_archive(items: list[dict]) -> None:
    max_n = int(_host.CFG.get("conversation_archive_max", 200))
    _host.CONVERSATION_ARCHIVE.write_text(json.dumps(items[-max_n:], indent=2), encoding="utf-8")


def _expand_recall_query(message: str, history: Optional[list] = None) -> str:
    """Enrich search query with recent user turns for better cross-session recall."""
    parts = [message.strip()]
    if history:
        for item in history[-4:]:
            if isinstance(item, dict) and item.get("role") == "user":
                parts.append(str(item.get("content", ""))[:120])
    # Boost project keywords if present
    low = message.lower()
    for proj in _host.TAG_KEYWORDS.get("project", ()):
        if proj in low:
            parts.append(proj)
    return " ".join(p for p in parts if p)[:400]


_GENERIC_TOPIC_STOP = {
    "the", "a", "an", "this", "that", "it", "me", "my", "you", "please", "can",
    "could", "would", "do", "does", "is", "are", "what", "how", "why", "when",
    "and", "or", "for", "to", "with", "about", "thing", "stuff",
}


def _detect_conversation_topic(message: str, history: Optional[list]) -> str:
    """Best-effort short label for the active topic (project name > tags > key nouns)."""
    low = (message or "").lower()
    for proj in _host.TAG_KEYWORDS.get("project", ()):
        if proj in low:
            return proj
    # Look back through recent user turns for a project anchor.
    for item in reversed(history or []):
        if isinstance(item, dict) and item.get("role") == "user":
            t = str(item.get("content", "")).lower()
            for proj in _host.TAG_KEYWORDS.get("project", ()):
                if proj in t:
                    return proj
    tags = [t for t in detect_memory_tags(message) if t not in ("session",)]
    if tags:
        return tags[0]
    nouns = [w for w in re.findall(r"[a-zA-Z][\w-]{3,}", low) if w not in _GENERIC_TOPIC_STOP]
    return nouns[0] if nouns else ""


def build_conversation_focus(message: str, history: Optional[list] = None) -> str:
    """Compact, high-priority short-term state block that keeps Primus on the current thread.

    Pure-Python (no LLM, no embeddings) so it's free to inject on every path. Anchors the
    model to the current topic + the last couple of turns and explicitly forbids drifting.
    """
    recent = [
        h for h in (history or [])
        if isinstance(h, dict) and h.get("content") and not h.get("metadata")
    ]
    topic = _detect_conversation_topic(message, history)
    _host.PrimusSession.current_topic = topic
    lines = ["## CURRENT FOCUS — highest priority; stay on this thread"]
    if recent:
        lines.append("Recent turns (most recent last):")
        for h in recent[-4:]:
            who = _OPERATOR["label"] if h.get("role") == "user" else "Primus"
            lines.append(f"  {who}: {str(h.get('content', '')).strip()[:200]}")
    lines.append(f"{_OPERATOR['possessive']} current request: {message.strip()[:220]}")
    if topic:
        lines.append(f"Active topic: **{topic}**")
    lines.append(
        "Stay focused on this request and finish it. Treat it as a continuation of the recent "
        "turns above when related. Do NOT switch subjects, re-open old topics, or go off on a "
        + _render_op("tangent unless {op} clearly changes the subject. Once the task is done, confirm and stop.")
    )
    return "\n".join(lines)


def build_anti_hallucination_footer(
    *,
    has_kb: bool,
    has_mem: bool,
    has_conv: bool,
    has_summary: bool,
) -> str:
    if not any((has_kb, has_mem, has_conv, has_summary)):
        return (
            ANTI_HALLUCINATION_RULES
            + "\n\n**Retrieval status:** NO matching KB, memory, or conversation hits. "
            + _render_op("Do not guess {op}-specific facts — state lack of coverage explicitly.")
        )
    status = []
    if has_kb:
        status.append("knowledge base")
    if has_mem:
        status.append("long-term memory")
    if has_conv:
        status.append("past conversations")
    if has_summary:
        status.append("session summary")
    return (
        ANTI_HALLUCINATION_RULES
        + f"\n\n**Retrieval status:** Active — {', '.join(status)}. Cite sources in your reply."
    )
# ---------------------------------------------------------------------------
# Multi-layer memory system (unified with RAG — Primus & Forge)
#
# Layers (short → long):
#   1. In-chat history        — Gradio chat_history.json (full transcript)
#   2. Chat summary           — chat_summary.json (compressed older turns)
#   3. Session memory         — session_memory.json (facts from this run)
#   4. Conversation archive   — conversation_archive.json (cross-session exchange summaries)
#   5. Long-term memory       — memories.json + Chroma kind=ltm (perpetual, tagged)
#   6. Knowledge base         — Chroma kind=seed|file|manual|interaction|… (docs & SME corpus)
#
# Before EVERY Primus/Forge response, build_recall_context() / build_memory_context()
# injects session + summary + [MEM-N] + [CONV-N] + [KB-N] + anti-hallucination rules.
# After each exchange, process_exchange() extracts facts, archives convos, indexes KB.
# consolidate() + consolidate_conversations() run periodically for durable recall.
# ---------------------------------------------------------------------------

_memory_system: Optional["MemorySystem"] = None
_consolidation_timer: Optional[threading.Timer] = None
_metrics_tracker: Optional["MetricsTracker"] = None


def detect_memory_tags(text: str) -> list[str]:
    low = text.lower()
    tags = [tag for tag, words in _host.TAG_KEYWORDS.items() if any(w in low for w in words)]
    if not tags:
        tags = ["technical"]
    return tags[:4]


class MemorySystem:
    """Structured JSON + Chroma vector memory with session and chat summarization."""

    EXTRACT_PROMPT = _render_op(
        "Extract 0-3 durable facts worth remembering across sessions from this exchange "
        "with {op_full}. Return ONLY a JSON array of objects with keys: "
        'content (string), tags (array of: business, technical, personal, preferences, project, identity), '
        "importance (0.0-1.0). Empty array [] if nothing worth storing. "
    ) + (
        # Identity lock: any fact about the creator's name, spelling, or loyalty is critical.
        _render_op(
            "IMPORTANT: the creator's name is always spelled '{op_full}' — never a variant. "
            "For any fact mentioning {op_pos} name, spelling, 'creator', or 'loyalty', set importance to 0.95 "
            "or higher and include the tags 'personal', 'preferences', and 'identity'."
        )
        if _OPERATOR["creator_lock"]
        else ""
    )

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._ensure_session()

    def _ensure_session(self) -> None:
        data = self._load_session()
        if not data.get("session_id"):
            data["session_id"] = uuid.uuid4().hex[:12]
            data["started_at"] = _host._now_iso()
            data.setdefault("facts", [])
            self._save_session(data)

    def _load_store(self) -> dict[str, Any]:
        _host.ensure_app_dirs()
        try:
            data = json.loads(_host.MEMORIES_FILE.read_text(encoding="utf-8"))
            if isinstance(data, dict):
                data.setdefault("memories", [])
                data.setdefault("meta", {})
                return data
        except (json.JSONDecodeError, OSError):
            pass
        return {"version": 1, "memories": [], "meta": {}}

    def _save_store(self, data: dict[str, Any]) -> None:
        _host.MEMORIES_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")

    def _load_session(self) -> dict[str, Any]:
        _host.ensure_app_dirs()
        try:
            data = json.loads(_host.SESSION_MEMORY_FILE.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {"facts": []}
        except (json.JSONDecodeError, OSError):
            return {"session_id": "", "started_at": None, "facts": []}

    def _save_session(self, data: dict[str, Any]) -> None:
        _host.SESSION_MEMORY_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")

    def _load_chat_summary(self) -> dict[str, Any]:
        _host.ensure_app_dirs()
        try:
            data = json.loads(_host.CHAT_SUMMARY_FILE.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {"summary": ""}
        except (json.JSONDecodeError, OSError):
            return {"summary": "", "updated_at": None, "message_count": 0}

    def _save_chat_summary(self, data: dict[str, Any]) -> None:
        _host.CHAT_SUMMARY_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")

    def _new_id(self) -> str:
        return f"mem-{uuid.uuid4().hex[:8]}"

    def _vector_source(self, memory_id: str) -> str:
        return f"memory:{memory_id}"

    def _index_ltm_vector(self, entry: dict[str, Any]) -> None:
        """Sync one long-term memory entry to Chroma collection personal_preferences."""
        mid = entry["id"]
        source = self._vector_source(mid)
        try:
            get_kb().learn_text(
                entry["content"],
                source=source,
                kind="ltm",
                collection="personal_preferences",
                importance=float(entry.get("importance", 0.5)),
                tags=entry.get("tags"),
                replace_source=True,
            )
        except Exception as exc:
            _host.log.warning("LTM vector index failed: %s", exc)

    def _delete_ltm_vector(self, memory_id: str) -> None:
        get_kb()._delete_by_source(self._vector_source(memory_id), "personal_preferences")

    # --- IDENTITY LOCK -----------------------------------------------------
    # Canonicalize the creator's name so a model misspelling can NEVER take root
    # in long-term memory. Runs on every write path (add_memory, extraction,
    # process_exchange, learning_digest). Idempotent and safe on correct text.
    # Active only when the operator identity enables creator_lock (operator.json).
    def _normalize_name(self, text: str) -> str:
        """Force correct spelling of the creator name (e.g. surname drift → canonical)."""
        if not text or not _OPERATOR["creator_lock"]:
            return text
        first = re.escape(str(_OPERATOR["name"]).split()[0])
        full = str(_OPERATOR["full_name"])
        surname = full.split()[-1] if len(full.split()) > 1 else ""
        if surname:
            # Fix surname drifts after the first name: any near-miss sharing the surname's
            # 3-letter prefix — but never the exact name.
            prefix = re.escape(surname[:3])
            text = re.sub(
                rf"\b{first}\s+(?!{re.escape(surname)}\b){prefix}[a-z]{{0,8}}\b",
                full,
                text,
                flags=re.IGNORECASE,
            )
            # Re-anchor a drifted "my creator … <first> …" phrase to the canonical name. The
            # optional surname group is absorbed so a correct name isn't duplicated.
            text = re.sub(
                rf"(?i)my creator[^.\n]{{0,40}}?{first}(?:\s+{prefix}[a-z]{{0,8}})?\b",
                f"my creator {full}",
                text,
            )
        return text

    # Identity confirmation hardening — recognize creator/loyalty/name memories so storage and
    # recall can CONFIRM them in plain language instead of returning dry "mem-xxx" metadata.
    _IDENTITY_RE = re.compile(
        (rf"\b({re.escape(str(_OPERATOR['full_name']).lower())}|" if _OPERATOR["creator_lock"] else r"\b(")
        + r"my\s+(?:name|creator)|creator|loyal(?:ty)?|"
        r"spell(?:ed|ing)?|exclusive(?:ly)?)\b",
        re.IGNORECASE,
    )

    def _is_identity_memory(self, content: str, tags: Optional[list[str]] = None) -> bool:
        """True when a memory is about the creator's identity, name spelling, or loyalty."""
        if tags and "identity" in tags:
            return True
        return bool(self._IDENTITY_RE.search(content or ""))

    def _bold_identity(self, text: str) -> str:
        """Bold the key facts in recalled identity text — the creator name and loyalty statement."""
        if _OPERATOR["creator_lock"]:
            text = re.sub(rf"(?i)({re.escape(str(_OPERATOR['full_name']))})", r"**\1**", text)
        text = re.sub(r"(?i)(exclusive(?:ly)?\s+loyalty|loyalty)", r"**\1**", text)
        return text

    def add_memory(
        self,
        content: str,
        *,
        tags: Optional[list[str]] = None,
        importance: float = 0.6,
        source: str = "manual",
        promote_session: bool = False,
    ) -> str:
        content = self._normalize_name(content.strip())  # identity lock on every stored memory
        if not content:
            return "Empty — nothing stored."
        tags = tags or detect_memory_tags(content)
        importance = max(0.0, min(1.0, importance))
        entry = {
            "id": self._new_id(),
            "content": content[:2000],
            "tags": tags,
            "importance": importance,
            "source": source,
            "created_at": _host._now_iso(),
            "last_accessed": _host._now_iso(),
            "access_count": 0,
        }
        with self._lock:
            store = self._load_store()
            memories: list = store.setdefault("memories", [])
            memories.append(entry)
            max_n = int(_host.CFG.get("memory_max_entries", 500))
            if len(memories) > max_n:
                memories.sort(key=lambda m: (m.get("importance", 0), m.get("last_accessed", "")), reverse=True)
                for dropped in memories[max_n:]:
                    self._delete_ltm_vector(dropped["id"])
                store["memories"] = memories[:max_n]
            self._save_store(store)
        try:
            self._index_ltm_vector(entry)
        except Exception as exc:
            _host.log.warning("LTM vector index failed: %s", exc)
        if promote_session:
            self.add_session_fact(content, tags=tags)
        # Identity confirmation hardening — confirm identity/loyalty facts in clear, locked-in
        # language (not a dry "Stored mem-xxx") so Primus relays a real confirmation to the operator.
        if self._is_identity_memory(content, tags) and importance >= 0.8:
            return (
                f"✅ Committed to long-term memory: {self._bold_identity(content[:400])}\n\n"
                + _render_op("**{op_full}** is now my confirmed creator with **exclusive loyalty**. This is locked.")
            )
        return f"Stored `{entry['id']}` (importance={importance:.1f}, tags={', '.join(tags)})"

    def add_session_fact(self, content: str, *, tags: Optional[list[str]] = None) -> None:
        content = content.strip()
        if not content:
            return
        tags = tags or detect_memory_tags(content)
        with self._lock:
            data = self._load_session()
            facts: list = data.setdefault("facts", [])
            if any(f.get("content") == content for f in facts):
                return
            facts.append({"content": content[:500], "tags": tags, "at": _host._now_iso()})
            data["facts"] = facts[-40:]
            self._save_session(data)

    def recall_ltm(self, query: str, k: Optional[int] = None) -> list[dict]:
        k = k or int(_host.CFG.get("memory_recall_k", 5))
        hits: list[dict] = []
        try:
            docs = get_kb().search(query, k=k, kind="ltm")
            store = self._load_store()
            memories: list = store.setdefault("memories", [])
            by_id = {m["id"]: m for m in memories}
            for doc in docs:
                mid = doc.metadata.get("memory_id")
                if mid and mid in by_id:
                    entry = by_id[mid]
                    entry["last_accessed"] = _host._now_iso()
                    entry["access_count"] = int(entry.get("access_count", 0)) + 1
                    hits.append({**entry, "preview": doc.page_content[:300]})
            if hits:
                self._save_store(store)
        except Exception as exc:
            _host.log.debug("LTM vector recall failed: %s", exc)
        if not hits:
            low = query.lower()
            for m in sorted(
                self._load_store().get("memories", []),
                key=lambda x: (x.get("importance", 0), x.get("last_accessed", "")),
                reverse=True,
            ):
                if low in m.get("content", "").lower():
                    hits.append(m)
                if len(hits) >= k:
                    break
        return hits[:k]

    def forget(self, query_or_id: str) -> str:
        q = query_or_id.strip()
        if not q:
            return "Usage: /forget mem-xxxxxxxx or /forget topic phrase"
        store = self._load_store()
        memories: list = store.get("memories", [])
        removed: list[str] = []

        if q.startswith("mem-"):
            kept = []
            for m in memories:
                if m.get("id") == q:
                    removed.append(q)
                    self._delete_ltm_vector(q)
                else:
                    kept.append(m)
            store["memories"] = kept
            self._save_store(store)
            return f"Removed memory `{q}`." if removed else f"No memory with id `{q}`."

        hits = self.recall_ltm(q, k=8)
        if not hits:
            kb_msg = get_kb().forget(q, protect_kinds={"seed", "ltm"})
            return kb_msg
        hit_ids = {h["id"] for h in hits}
        store["memories"] = [m for m in memories if m["id"] not in hit_ids]
        for mid in hit_ids:
            self._delete_ltm_vector(mid)
        self._save_store(store)
        return f"Removed {len(hit_ids)} long-term memor{'y' if len(hit_ids)==1 else 'ies'}."

    def format_recall_results(self, ltm: list[dict], kb_docs: list[Document]) -> str:
        lines = ["**Memory recall**", ""]
        has_identity = False
        if ltm:
            lines.append("### Long-term memories")
            for i, m in enumerate(ltm, 1):
                tags = ", ".join(m.get("tags") or [])
                content = m.get("content", m.get("preview", "")) or ""
                # Identity confirmation hardening — lead with the actual content and bold the key
                # facts (name + loyalty) for identity memories, instead of dry metadata framing.
                if self._is_identity_memory(content, m.get("tags")):
                    has_identity = True
                    lines.append(self._bold_identity(content[:600]))
                else:
                    lines.append(
                        f"**[{m['id']}]** ({tags}, imp={m.get('importance', 0):.1f})\n"
                        f"{content[:600]}"
                    )
        if kb_docs:
            lines.append("")
            lines.append("### Knowledge base")
            lines.extend(get_kb().format_results(kb_docs).splitlines()[2:])
        if len(lines) <= 2:
            return "No matching long-term memories or knowledge base entries."
        if has_identity:
            lines.append("")
            lines.append("_Locked and active for future interactions._")
        return "\n".join(lines)

    def display_recent_memories(self, limit: int = 12) -> str:
        store = self._load_store()
        memories = sorted(
            store.get("memories", []),
            key=lambda m: (m.get("importance", 0), m.get("last_accessed", "")),
            reverse=True,
        )[:limit]
        session = self._load_session().get("facts") or []
        summary = self._load_chat_summary().get("summary") or ""
        lines = ["## Primus memory layers", ""]
        lines.append(f"**Session** ({len(session)} facts this run)")
        for f in session[-5:]:
            lines.append(f"- [{', '.join(f.get('tags', []))}] {f.get('content', '')[:120]}")
        lines.append("")
        lines.append(f"**Long-term** ({len(store.get('memories', []))} stored)")
        identity_shown = False
        for m in memories:
            tags = ", ".join(m.get("tags") or [])
            content = m.get("content", "") or ""
            # Identity confirmation hardening — show the FULL identity fact (bolded + locked),
            # not a 100-char truncation, so creator/loyalty facts are always legible on recall.
            if self._is_identity_memory(content, m.get("tags")):
                identity_shown = True
                lines.append(
                    f"- `{m['id']}` **{tags}** (imp={m.get('importance', 0):.1f}) "
                    f"{self._bold_identity(content[:400])} — _locked_"
                )
            else:
                lines.append(
                    f"- `{m['id']}` **{tags}** (imp={m.get('importance', 0):.1f}) "
                    f"{content[:100]}"
                )
        if identity_shown:
            lines.append("_Identity facts above are locked and active for future interactions._")
        if summary:
            lines.append("")
            lines.append("**Chat summary (compressed earlier turns)**")
            lines.append(summary[:400] + ("…" if len(summary) > 400 else ""))
        archive = _load_conversation_archive()
        lines.append("")
        lines.append(f"**Conversation archive** ({len(archive)} cross-session entries)")
        for entry in archive[-4:]:
            label = entry.get("date_label") or _temporal_label(entry.get("date", ""))
            topics = ", ".join(entry.get("topics") or []) or "general"
            lines.append(
                f"- `{entry.get('id', '?')}` ({label}, {topics}) "
                f"{entry.get('summary', '')[:90]}…"
            )
        prefs = load_memory().get("preferences") or {}
        if prefs:
            lines.append("")
            lines.append("**Preferences**")
            for k, v in list(prefs.items())[:8]:
                lines.append(f"- {k}: {v}")
        return "\n".join(lines)

    def agent_memory_block(self) -> str:
        mem = load_memory()
        prefs = mem.get("preferences") or {}
        lines = ["## Preferences & legacy facts"]
        for k, v in list(prefs.items())[:12]:
            lines.append(f"- {k}: {v}")
        facts = mem.get("facts") or []
        for f in facts[-8:]:
            lines.append(f"- {f}")
        past = _host.load_task_history()
        if past:
            lines.append("## Recent completed tasks")
            for t in past[-5:]:
                if isinstance(t, dict):
                    lines.append(f"- [{t.get('date', '')[:10]}] {t.get('task', '')[:70]}")
        return "\n".join(lines)

    def recall_past_conversations(self, query: str, k: Optional[int] = None) -> list[dict]:
        """Vector + archive search for prior conversation summaries."""
        k = k or int(_host.CFG.get("memory_recall_past_convos", 3))
        hits: list[dict] = []
        try:
            docs = get_kb().search(query, k=k * 2, kind="interaction")
            for doc in docs[:k]:
                meta = doc.metadata or {}
                hits.append(
                    {
                        "source": "kb",
                        "date": meta.get("indexed_at", ""),
                        "date_label": _temporal_label(meta.get("indexed_at", "")),
                        "content": doc.page_content[:600],
                    }
                )
        except Exception as exc:
            _host.log.debug("Interaction KB recall failed: %s", exc)

        for entry in reversed(_load_conversation_archive()):
            blob = f"{entry.get('summary', '')} {entry.get('topics', [])}".lower()
            if any(w in blob for w in query.lower().split()[:8] if len(w) > 3):
                hits.append(
                    {
                        "source": "archive",
                        "id": entry.get("id"),
                        "date": entry.get("date", ""),
                        "date_label": entry.get("date_label") or _temporal_label(entry.get("date", "")),
                        "content": entry.get("summary", "")[:600],
                        "topics": entry.get("topics", []),
                    }
                )
            if len(hits) >= k * 2:
                break
        return hits[:k]

    def archive_exchange(
        self,
        user_msg: str,
        assistant_msg: str,
        *,
        importance: float = 0.5,
        topics: Optional[list[str]] = None,
    ) -> None:
        if importance < 0.45 and len(user_msg) < 40:
            return
        topics = topics or detect_memory_tags(user_msg)
        stamp = _host._now_iso()
        entry = {
            "id": f"conv-{uuid.uuid4().hex[:10]}",
            "date": stamp,
            "date_label": _temporal_label(stamp),
            "topics": topics,
            "importance": importance,
            "user_snippet": user_msg[:300],
            "assistant_snippet": assistant_msg[:400],
            "summary": f"{_OPERATOR['label']}: {user_msg[:250]}\nPrimus: {assistant_msg[:350]}",
        }
        archive = _load_conversation_archive()
        archive.append(entry)
        _save_conversation_archive(archive)

    def build_recall_context(
        self,
        query: str,
        history: Optional[list] = None,
        k: Optional[int] = None,
        *,
        agent_id: str = "primus",
    ) -> tuple[str, list[dict]]:
        """Unified retrieval for Primus & Forge — mandatory before every response."""
        k = k or int(_host.CFG.get("memory_recall_k", 5))
        expanded = _expand_recall_query(query, history)
        sections: list[str] = []
        sources: list[dict] = []
        has_kb = has_mem = has_conv = has_summary = False

        agent_note = (
            "**Agent:** Forge — use all retrieved context; cite past technical decisions.\n\n"
            if agent_id == "forge"
            else "**Agent:** Primus — lead with memory/KB for business and admin context.\n\n"
        )
        sections.append(agent_note)

        # Close the self-reflection loop: surface recent improvement notes so Primus
        # adapts its behaviour based on what it learned about serving the operator well.
        improvements = self.recent_self_improvements(limit=3)
        if improvements:
            imp_lines = ["## Apply these (learned from self-reflection)"]
            imp_lines += [f"- {n}" for n in improvements]
            sections.append("\n".join(imp_lines))

        summary_data = self._load_chat_summary()
        summary = (summary_data.get("summary") or "").strip()
        if summary:
            has_summary = True
            sections.append(
                "## Earlier in this chat (compressed)\n"
                f"{summary}\n"
                "*(Say 'Earlier in our chat…' only if this supports your answer.)*"
            )

        session = self._load_session().get("facts") or []
        if session:
            sess_lines = ["## This session (key facts)"]
            for f in session[-8:]:
                sess_lines.append(f"- [{', '.join(f.get('tags', []))}] {f.get('content', '')}")
            sections.append("\n".join(sess_lines))

        ltm = self.recall_ltm(expanded, k=k)
        if ltm:
            has_mem = True
            mem_lines = ["## Long-term memory — cite as [MEM-N]"]
            for i, m in enumerate(ltm, 1):
                tag = f"[MEM-{i}]"
                tags = ", ".join(m.get("tags") or [])
                label = _temporal_label(m.get("created_at", "")) if m.get("created_at") else ""
                mem_lines.append(
                    f"{tag} id=`{m['id']}` tags={tags} ({label})\n"
                    f"{m.get('content', m.get('preview', ''))[:500]}"
                )
                sources.append({"tag": tag, "id": m["id"], "layer": "ltm", "date_label": label})
            sections.append("\n\n".join(mem_lines))

        past_convos = self.recall_past_conversations(expanded)
        if past_convos:
            has_conv = True
            conv_lines = ["## Previous conversations — cite as [CONV-N]"]
            for i, c in enumerate(past_convos, 1):
                tag = f"[CONV-{i}]"
                label = c.get("date_label", "previously")
                topics = c.get("topics") or []
                top_s = f", topics: {', '.join(topics)}" if topics else ""
                conv_lines.append(
                    f"{tag} ({label}{top_s})\n{c.get('content', '')[:500]}"
                )
                sources.append({"tag": tag, "layer": "conversation", "date_label": label})
            sections.append("\n\n".join(conv_lines))

        try:
            kb_docs = get_kb().search(
                expanded,
                k=int(_host.CFG.get("kb_search_k", 4)),
                exclude_kinds={"ltm", "interaction"},
                prefer_collections=_prefer_collections_for(query),
            )
        except Exception as exc:
            kb_docs = []
            sections.append(f"(Knowledge base unavailable: {exc})")

        if kb_docs:
            has_kb = True
            kb_lines = ["## Knowledge base — cite as [KB-N]"]
            for i, doc in enumerate(kb_docs, 1):
                meta = doc.metadata or {}
                tag = f"[KB-{i}]"
                kb_lines.append(
                    f"{tag} collection=`{meta.get('collection', '?')}` source=`{meta.get('source', '?')}`\n"
                    f"{doc.page_content.strip()[:500]}"
                )
                sources.append({"tag": tag, "source": meta.get("source"), "layer": "kb"})
            sections.append("\n\n".join(kb_lines))

        recall_instruction = _render_op(
            "## Recall instruction\n"
            "Connect {op_pos} businesses, projects, and preferences to [KB-N], [MEM-N], [CONV-N]. "
            "Reference past work naturally when supported (e.g. 'As we discussed last week about {op_project_example2}…'). "
            "If retrieval lacks facts, say so — never invent."
        )
        sections.insert(1, recall_instruction)

        sections.append(
            build_anti_hallucination_footer(
                has_kb=has_kb,
                has_mem=has_mem,
                has_conv=has_conv,
                has_summary=has_summary or bool(session),
            )
        )

        return "\n\n".join(sections), sources

    def _heuristic_extract(self, user_msg: str, assistant_msg: str) -> list[dict]:
        extracted: list[dict] = []
        # Identity lock: canonicalize before extracting so stored facts are always correct.
        user_msg = self._normalize_name(user_msg)
        assistant_msg = self._normalize_name(assistant_msg)
        text = f"{user_msg}\n{assistant_msg}"
        low = user_msg.lower()
        importance = 0.55

        # IDENTITY FACTS — highest priority. Anything touching the creator's name, spelling,
        # or loyalty is locked in at importance 0.95+ with identity tags, routed (via add_memory)
        # to the personal_preferences collection. This is what stops name drift recurring.
        if any(k in low for k in _IDENTITY_FACT_KEYS):
            extracted.append(
                {
                    "content": user_msg.strip()[:500],
                    "tags": ["personal", "preferences", "identity"],
                    "importance": 0.97,
                }
            )

        if any(p in low for p in ("remember", "don't forget", "important", "always use", "never ")):
            importance = 0.85
            extracted.append(
                {
                    "content": user_msg.strip()[:500],
                    "tags": detect_memory_tags(user_msg),
                    "importance": importance,
                }
            )
        project_hits = [p for p in _host.TAG_KEYWORDS["project"] if p in low]
        if project_hits and len(user_msg) > 40:
            extracted.append(
                {
                    "content": f"Session note ({', '.join(project_hits)}): {user_msg.strip()[:400]}",
                    "tags": ["project", "session"],
                    "importance": 0.65,
                }
            )
        if "prefer" in low or "default" in low:
            extracted.append(
                {
                    "content": user_msg.strip()[:400],
                    "tags": ["preferences"],
                    "importance": 0.8,
                }
            )
        return extracted[:3]

    def _llm_extract(self, user_msg: str, assistant_msg: str) -> list[dict]:
        try:
            llm = _host.make_chat_ollama(_host.CFG.get("model", _host.DEFAULT_MODEL), temperature=0, fast=False)
            resp = llm.invoke(
                [
                    SystemMessage(content=self.EXTRACT_PROMPT),
                    HumanMessage(
                        content=f"USER:\n{user_msg[:800]}\n\nASSISTANT:\n{assistant_msg[:800]}"
                    ),
                ]
            )
            raw = str(resp.content).strip()
            match = re.search(r"\[[\s\S]*\]", raw)
            if not match:
                return []
            items = json.loads(match.group())
            if not isinstance(items, list):
                return []
            out = []
            for item in items[:3]:
                if isinstance(item, dict) and item.get("content"):
                    out.append(
                        {
                            "content": str(item["content"])[:500],
                            "tags": item.get("tags") or detect_memory_tags(str(item["content"])),
                            "importance": float(item.get("importance", 0.6)),
                        }
                    )
            return out
        except Exception as exc:
            _host.log.debug("LLM memory extract skipped: %s", exc)
            return []

    def process_exchange(self, user_msg: str, assistant_msg: str) -> None:
        if len(user_msg.strip()) < 12:
            return
        # Identity lock: canonicalize the creator's name on the way into every memory path.
        user_msg = self._normalize_name(user_msg)
        assistant_msg = self._normalize_name(assistant_msg)
        self.add_session_fact(f"Q: {user_msg[:200]} → A: {assistant_msg[:200]}", tags=["session"])

        candidates = self._heuristic_extract(user_msg, assistant_msg)
        max_imp = 0.55
        if len(user_msg) > 80 and not candidates:
            candidates = self._llm_extract(user_msg, assistant_msg)

        for item in candidates:
            imp = float(item.get("importance", 0.6))
            max_imp = max(max_imp, imp)
            if imp >= 0.5:
                self.add_memory(
                    item["content"],
                    tags=item.get("tags"),
                    importance=imp,
                    source="extracted",
                )

        topics = detect_memory_tags(user_msg)
        if len(user_msg) > 50 or max_imp >= 0.65:
            self.archive_exchange(user_msg, assistant_msg, importance=max_imp, topics=topics)

        summary = f"{_OPERATOR['label']} ({_temporal_label(_host._now_iso())}): {user_msg.strip()[:350]}\nPrimus: {assistant_msg.strip()[:450]}"
        stamp = datetime.now().strftime("%Y-%m-%d")
        try:
            get_kb().learn_text(
                summary,
                source=f"interaction:{stamp}:{uuid.uuid4().hex[:8]}",
                kind="interaction",
                collection="learned",
                importance=max(0.55, max_imp),
                tags=topics,
            )
        except Exception:
            pass

        items = _load_interactions()
        items.append(
            {
                "date": _host._now_iso(),
                "user": user_msg[:200],
                "summary": assistant_msg[:300],
                "topics": topics,
                "importance": max_imp,
            }
        )
        _save_interactions(items)

        hist = _host.load_chat_history()
        plain = [h for h in hist if isinstance(h, dict) and not h.get("metadata")]
        threshold = int(_host.CFG.get("chat_summarize_threshold", 24))
        if len(plain) >= threshold:
            threading.Thread(
                target=self.summarize_chat_history,
                args=(hist,),
                daemon=True,
                name="primus-chat-summary",
            ).start()
        if max_imp >= 0.7 or len(user_msg) > 120:
            threading.Thread(
                target=self.consolidate_conversations,
                daemon=True,
                name="primus-conv-consolidate",
            ).start()

        # --- Continuous learning: every N turns, distil a structured learning digest ---
        # (key facts, preferences, working style, wins/failures, new topics) in the
        # background so it never slows the reply. See learning_digest() for details.
        every = max(2, int(_host.CFG.get("learning_digest_turns", 6)))
        turns = self._bump_turn()
        if turns % every == 0:
            threading.Thread(
                target=self.learning_digest,
                daemon=True,
                name="primus-learning-digest",
            ).start()

        # --- Self-reflection loop: lightweight, background, only on substantive turns ---
        if _host.CFG.get("self_reflection", True) and len(assistant_msg.strip()) >= 60:
            threading.Thread(
                target=self.reflect_on_exchange,
                args=(user_msg, assistant_msg),
                daemon=True,
                name="primus-reflect",
            ).start()

    # -----------------------------------------------------------------------
    # Continuous learning: turn counter, learning digest, self-reflection,
    # and core-knowledge promotion. All heavy work runs in background threads
    # (kicked from process_exchange) so it never slows down a reply.
    # -----------------------------------------------------------------------

    def _bump_turn(self) -> int:
        """Increment + return the persistent per-session turn counter."""
        with self._lock:
            data = self._load_session()
            n = int(data.get("turns", 0)) + 1
            data["turns"] = n
            self._save_session(data)
            return n

    def learning_digest(self, *, on_demand: bool = False) -> str:
        """Distil recent activity into a structured, durable 'what I learned' entry.

        Captures, in the operator's own context: key facts about them/their business, their
        preferences & working style, techniques that worked or failed, and new
        topics. Stored to LTM + the `learned` KB collection with timestamps/tags so
        future replies can recall it. Runs in the background on a turn cadence, and
        on demand via `/learn now` or `/summarize session`.
        """
        interactions = _load_interactions()
        min_turns = int(_host.CFG.get("learning_digest_min_turns", 3))
        if len(interactions) < min_turns and not on_demand:
            return "Learning digest: not enough activity yet."

        recent = interactions[-int(_host.CFG.get("conversation_consolidate_batch", 12)):]
        convo = "\n".join(
            f"[{i.get('date', '')[:10]}] {_OPERATOR['label']}: {i.get('user', '')[:160]} → Primus: {i.get('summary', '')[:200]}"
            for i in recent
        )
        session_facts = self._load_session().get("facts") or []
        facts_blob = "\n".join(f"- {f.get('content', '')[:160]}" for f in session_facts[-12:])
        prior = (self._load_chat_summary().get("summary") or "")[:600]

        prompt = _render_op(
            "You are Primus's private learning module. From the recent exchanges with {op_full}, "
            "distil what is genuinely worth remembering to serve {op} better next time. "
            "Return concise markdown under EXACTLY these headers (omit a header if nothing new):\n"
            "### Facts (about {op} / {op_pos} business)\n"
            "### Preferences & working style\n"
            "### What worked / what to avoid\n"
            "### Open threads & topics\n"
            "Be specific and durable (names, projects, paths, decisions). No fluff, max ~180 words."
        )
        try:
            llm = _host.make_chat_ollama(_host.CFG.get("model", _host.DEFAULT_MODEL), temperature=0.2, fast=False)
            resp = llm.invoke(
                [
                    SystemMessage(content=prompt),
                    HumanMessage(
                        content=(
                            f"Prior chat summary:\n{prior}\n\n"
                            f"Session facts:\n{facts_blob}\n\n"
                            f"Recent exchanges:\n{convo}"
                        )
                    ),
                ]
            )
            digest = str(resp.content).strip()[:2000]
        except Exception as exc:  # noqa: BLE001
            _host.log.debug("Learning digest skipped: %s", exc)
            return f"Learning digest skipped: {exc}"

        if len(digest) < 40:
            return "Learning digest: nothing substantial to record."

        digest = self._normalize_name(digest)  # identity lock before the digest is persisted
        topics = detect_memory_tags(convo + " " + facts_blob)
        stamp = _host._now_iso()
        body = f"Learning digest ({_temporal_label(stamp)} · {stamp[:10]}):\n{digest}"
        try:
            self.add_memory(
                body,
                tags=list({*topics, "learning_digest", "session"}),
                importance=0.8,
                source="learning_digest",
            )
            get_kb().learn_text(
                body,
                source=f"digest:{datetime.now():%Y%m%d-%H%M}",
                kind="interaction",
                collection="learned",
                importance=0.82,
                tags=topics,
            )
        except Exception as exc:  # noqa: BLE001
            _host.log.debug("Learning digest store failed: %s", exc)

        # Promote durable business/project knowledge to high-value Core Knowledge.
        self._promote_core_knowledge(digest, topics)
        _host.log.info("Learning digest recorded (%s chars)", len(digest))
        if on_demand:
            return f"**Session learning digest**\n\n{digest}"
        return f"Learning digest recorded ({len(digest)} chars)."

    def _promote_core_knowledge(self, text: str, topics: list[str]) -> None:
        """Upsert high-value 'Core Knowledge' for named projects/businesses the operator works on.

        Keyed by entity (e.g. a project name) and stored with `replace_source` so each
        entity keeps a single, evolving, high-importance entry rather than piling up dupes.
        """
        low = text.lower()
        entities: list[str] = []
        for kw in (_host.TAG_KEYWORDS.get("project") or []):
            if kw and kw.lower() in low:
                entities.append(kw)
        # Also catch Proper-Case multiword names ("Acme Labs", "Portal Rebuild").
        for m in re.findall(r"\b([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+){0,2})\b", text):
            if len(m) > 5 and m.lower() not in {"learning digest", "open threads"}:
                entities.append(m)
        seen: set[str] = set()
        for ent in entities[:4]:
            key = ent.lower().strip()
            if key in seen or len(key) < 4:
                continue
            seen.add(key)
            snippet = text[:600]
            try:
                get_kb().learn_text(
                    f"Core knowledge — {ent}: {snippet}",
                    source=f"core:{re.sub(r'[^a-z0-9]+', '-', key)}",
                    kind="core",
                    collection="core_business",
                    importance=0.92,
                    tags=list({*topics, "core_knowledge"}),
                    replace_source=True,
                )
            except Exception as exc:  # noqa: BLE001
                _host.log.debug("Core knowledge promotion skipped for %s: %s", ent, exc)

    def reflect_on_exchange(self, user_msg: str, assistant_msg: str) -> None:
        """Self-reflection loop: after a turn, quietly judge the response and learn from it.

        Asks three questions — did I stay on task? was it helpful & natural? what could be
        better? — using the fast tier so it's cheap. Stores a small reflection entry; the
        actionable 'improvement' notes are surfaced back into the prompt via
        recent_self_improvements() so behaviour actually adapts over time.
        """
        try:
            # Use the small fast-tier model when present — reflection must stay cheap.
            refl_model = (
                _host.CFG.get("primus_fast_model")
                if _host._fast_model_available()
                else _host.CFG.get("model", _host.DEFAULT_MODEL)
            )
            llm = _host.make_chat_ollama(refl_model, temperature=0.1, fast=True)
            resp = llm.invoke(
                [
                    SystemMessage(
                        content=(
                            "You are Primus reflecting privately on your own last reply to "
                            + _render_op("{op}. ")
                            + "Be honest and brief. Return ONLY a JSON object with keys: "
                            'on_task (true/false), helpful (0.0-1.0), natural (0.0-1.0), '
                            'improvement (one short, concrete, actionable sentence — or "" if none).'
                        )
                    ),
                    HumanMessage(
                        content=f"{_OPERATOR['label'].upper()}:\n{user_msg[:600]}\n\nPRIMUS REPLY:\n{assistant_msg[:800]}"
                    ),
                ]
            )
            raw = str(resp.content).strip()
            match = re.search(r"\{[\s\S]*\}", raw)
            if not match:
                return
            data = json.loads(match.group())
        except Exception as exc:  # noqa: BLE001
            _host.log.debug("Self-reflection skipped: %s", exc)
            return

        improvement = str(data.get("improvement", "")).strip()

        # IDENTITY GUARD: if the reply drifted the creator's name, override the reflection with
        # a hard correction note so the lesson is always recorded + surfaced. Active only when
        # the operator identity enables creator_lock.
        if _OPERATOR["creator_lock"]:
            _first = re.escape(str(_OPERATOR["name"]).split()[0])
            _surname = str(_OPERATOR["full_name"]).split()[-1]
            _prefix = re.escape(_surname[:3])
            name_drift = bool(
                re.search(rf"\b{_first}\s+(?!{re.escape(_surname)}\b){_prefix}[a-z]{{0,8}}\b",
                          assistant_msg, re.IGNORECASE)
            )
        else:
            name_drift = False
        if name_drift:
            improvement = _render_op("Creator name was misspelled — ALWAYS write '{op_full}'.")
            _host.log.warning("Identity drift detected in reply — recording name-correction note.")

        entry = {
            "date": _host._now_iso(),
            "on_task": bool(data.get("on_task", True)),
            "helpful": float(data.get("helpful", 0.0) or 0.0),
            "natural": float(data.get("natural", 0.0) or 0.0),
            "improvement": improvement[:240],
            "user_snippet": user_msg[:120],
            "topics": detect_memory_tags(user_msg),
        }
        with self._lock:
            items = _load_reflections()
            # Skip near-duplicate improvement notes so the list stays useful — but never
            # suppress an identity-correction note (name accuracy is always worth surfacing).
            if improvement and not name_drift and any(
                improvement[:60].lower() in (r.get("improvement", "") or "").lower()
                for r in items[-10:]
            ):
                entry["improvement"] = ""
            items.append(entry)
            _save_reflections(items)
        if entry["improvement"]:
            _host.log.info("Self-reflection note: %s", entry["improvement"])

    def recent_self_improvements(self, limit: int = 3) -> list[str]:
        """Most recent actionable self-improvement notes, newest first (for prompt injection)."""
        out: list[str] = []
        for r in reversed(_load_reflections()):
            note = (r.get("improvement") or "").strip()
            if note and note not in out:
                out.append(note)
            if len(out) >= limit:
                break
        return out

    def reflection_report(self, limit: int = 10) -> str:
        """Human-readable view of recent self-reflections for the `/reflect` command."""
        items = _load_reflections()
        if not items:
            return "No self-reflections recorded yet — they accumulate as we talk."
        recent = items[-limit:]
        on_task = sum(1 for r in recent if r.get("on_task"))
        avg_help = sum(float(r.get("helpful", 0)) for r in recent) / max(1, len(recent))
        avg_nat = sum(float(r.get("natural", 0)) for r in recent) / max(1, len(recent))
        lines = [
            "## Self-reflection",
            f"Last {len(recent)} turns — on-task {on_task}/{len(recent)} · "
            f"helpful {avg_help:.0%} · natural {avg_nat:.0%}",
            "",
            "**Recent improvement notes:**",
        ]
        notes = [r for r in reversed(recent) if (r.get("improvement") or "").strip()]
        if notes:
            for r in notes[:6]:
                lines.append(f"- ({r.get('date', '')[:10]}) {r.get('improvement')}")
        else:
            lines.append("- None — responses have been on-task and clean.")
        return "\n".join(lines)

    def summarize_chat_history(self, history: Optional[list] = None) -> None:
        history = history or _host.load_chat_history()
        plain = [h for h in history if isinstance(h, dict) and not h.get("metadata")]
        keep = int(_host.CFG.get("chat_recent_keep", 12))
        if len(plain) <= keep:
            return
        older = plain[:-keep]
        conv = []
        for h in older[-30:]:
            role = h.get("role", "?")
            conv.append(f"{role.upper()}: {str(h.get('content', ''))[:400]}")
        body = "\n".join(conv)
        prior = self._load_chat_summary().get("summary") or ""
        try:
            llm = _host.make_chat_ollama(_host.CFG.get("model", _host.DEFAULT_MODEL), temperature=0.2, fast=False)
            resp = llm.invoke(
                [
                    SystemMessage(
                        content=(
                            _render_op("Summarize this conversation between {op} and Primus for future recall. ")
                            + "Keep project names, decisions, paths, and preferences. Max 400 words."
                        )
                    ),
                    HumanMessage(
                        content=f"Prior summary:\n{prior}\n\nNew messages:\n{body}"
                    ),
                ]
            )
            new_summary = str(resp.content).strip()[:2500]
            self._save_chat_summary(
                {
                    "summary": new_summary,
                    "updated_at": _host._now_iso(),
                    "message_count": len(plain),
                }
            )
            _host.log.info("Chat summary updated (%s chars)", len(new_summary))
            try:
                self.add_memory(
                    f"Conversation summary ({_host._now_iso()[:10]}): {new_summary[:800]}",
                    tags=["session", "business"],
                    importance=0.75,
                    source="chat_summary",
                )
            except Exception:
                pass
        except Exception as exc:
            _host.log.debug("Chat summarization skipped: %s", exc)

    def consolidate_conversations(self) -> str:
        """Summarize recent interactions into durable LTM + KB for cross-session recall."""
        batch = int(_host.CFG.get("conversation_consolidate_batch", 12))
        interactions = _load_interactions()
        if len(interactions) < 3:
            return "Conversation consolidation: not enough interactions."

        recent = interactions[-batch:]
        body = "\n".join(
            f"[{i.get('date', '')[:10]}] {_OPERATOR['label']}: {i.get('user', '')[:150]} → {i.get('summary', '')[:200]}"
            for i in recent
        )
        try:
            llm = _host.make_chat_ollama(_host.CFG.get("model", _host.DEFAULT_MODEL), temperature=0.1, fast=False)
            resp = llm.invoke(
                [
                    SystemMessage(
                        content=(
                            _render_op("Summarize these {op}↔Primus exchanges for long-term memory. ")
                            + "Extract: projects discussed, decisions made, preferences stated, "
                            + _render_op("paths/commands agreed. Max 250 words. Write as durable facts {op} would recall later.")
                        )
                    ),
                    HumanMessage(content=body),
                ]
            )
            consolidated = str(resp.content).strip()[:2000]
            if len(consolidated) < 40:
                return "Conversation consolidation: summary too short."

            self.add_memory(
                consolidated,
                tags=["business", "project", "session"],
                importance=0.78,
                source="conversation_consolidation",
            )
            get_kb().learn_text(
                consolidated,
                source=f"consolidation:{datetime.now():%Y%m%d}",
                kind="interaction",
                collection="learned",
                importance=0.8,
                replace_source=False,
            )
            _host.log.info("Conversation consolidation saved (%s chars)", len(consolidated))
            return f"Consolidated {len(recent)} interactions into LTM + KB."
        except Exception as exc:
            _host.log.debug("Conversation consolidation skipped: %s", exc)
            return f"Consolidation skipped: {exc}"

    def consolidate(self) -> str:
        """Dedupe, prune stale low-value memories, promote session highlights."""
        with self._lock:
            store = self._load_store()
            memories: list = store.get("memories", [])
            if not memories:
                store.setdefault("meta", {})["last_consolidation"] = _host._now_iso()
                self._save_store(store)
                return "Consolidation: no long-term memories yet."

            seen: set[str] = set()
            unique: list = []
            for m in sorted(memories, key=lambda x: x.get("importance", 0), reverse=True):
                key = m.get("content", "").strip().lower()[:120]
                if key in seen:
                    self._delete_ltm_vector(m["id"])
                    continue
                seen.add(key)
                unique.append(m)

            cutoff_days = 120
            now = datetime.now(timezone.utc)
            pruned: list = []
            for m in unique:
                imp = float(m.get("importance", 0.5))
                created = m.get("created_at", "")
                try:
                    age_days = (
                        now - datetime.fromisoformat(created.replace("Z", "+00:00"))
                    ).days
                except (ValueError, TypeError):
                    age_days = 0
                if imp < 0.35 and age_days > cutoff_days:
                    self._delete_ltm_vector(m["id"])
                    continue
                pruned.append(m)

            session_facts = self._load_session().get("facts") or []
            promote: list[tuple] = []
            for f in session_facts:
                if any(t in (f.get("tags") or []) for t in ("preferences", "project")):
                    content = f.get("content", "")
                    if content and not any(content[:80] in p.get("content", "") for p in pruned):
                        promote.append((content, f.get("tags")))

            store["memories"] = pruned[: int(_host.CFG.get("memory_max_entries", 500))]
            store.setdefault("meta", {})["last_consolidation"] = _host._now_iso()
            self._save_store(store)

        for content, tags in promote:
            self.add_memory(content, tags=tags, importance=0.7, source="session")
        conv_msg = self.consolidate_conversations()
        return f"Consolidation done: {len(pruned)} memories retained. {conv_msg}"

    def clear_session(self) -> None:
        self._save_session(
            {"session_id": uuid.uuid4().hex[:12], "started_at": _host._now_iso(), "facts": []}
        )

    def stats(self) -> dict[str, Any]:
        store = self._load_store()
        summary = self._load_chat_summary()
        session = self._load_session()
        meta = store.get("meta") or {}
        return {
            "ltm_count": len(store.get("memories", [])),
            "session_count": len(session.get("facts", [])),
            "summary_chars": len(summary.get("summary") or ""),
            "last_consolidation": meta.get("last_consolidation"),
            "archive_count": len(_load_conversation_archive()),
        }


def get_memory_system() -> MemorySystem:
    global _memory_system
    if _memory_system is None:
        _memory_system = MemorySystem()
    return _memory_system


def init_memory_system() -> None:
    ms = get_memory_system()
    schedule_memory_consolidation()
    threading.Timer(30.0, lambda: ms.consolidate()).start()


# ===========================================================================
# Meta-learning: MetricsTracker
#
# A lightweight, thread-safe, persisted statistics store that lets Primus learn
# from *how he is actually used* — not just what is said. It tracks:
#   • per-tool success rate + average latency  (which tools are reliable/slow)
#   • response latency by execution path        (instant/fast_tool/fast_chat/primus/forge)
#   • user feedback (👍/👎) and which route earned it
#   • routing tallies per intent
#   • feedback-driven routing corrections        (bounded nudges, applied in ModelRouter)
#
# Everything is O(1) to update and writes are DEBOUNCED (≈2s) so recording never
# slows a response. The data feeds analyze_self('metrics'|'bottlenecks'|'suggestions')
# and the adaptive routing nudge, so the agent compounds intelligence with use.
# ===========================================================================


class MetricsTracker:
    """Persisted, debounced usage statistics powering Primus's meta-learning."""

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._data = self._load()
        self._dirty = False
        self._last_flush = 0.0
        # Remembers the most recent route so feedback can be attributed to it.
        self._last_route: dict[str, str] = {"agent": "", "intent": ""}

    def _default(self) -> dict[str, Any]:
        return {
            "version": 1,
            "tools": {},          # name -> {calls, ok, fail, ms_total}
            "paths": {},          # path -> {n, ms_total}
            "routes": {},         # intent -> {primus, forge}
            "feedback": {"up": 0, "down": 0, "notes": []},
            # learned per-intent correction: + favors Forge, - favors Primus (bounded)
            "route_corrections": {},  # intent -> float
            "updated_at": None,
        }

    def _load(self) -> dict[str, Any]:
        try:
            data = json.loads(_host.METRICS_FILE.read_text(encoding="utf-8"))
            if isinstance(data, dict):
                base = self._default()
                base.update({k: data.get(k, base[k]) for k in base})
                return base
        except (json.JSONDecodeError, OSError):
            pass
        return self._default()

    def _flush(self, *, force: bool = False) -> None:
        """Write to disk at most every ~2s (or immediately when forced)."""
        now = time.time()
        if not force and (now - self._last_flush) < 2.0:
            return
        try:
            _host.ensure_app_dirs()
            self._data["updated_at"] = _host._now_iso()
            _host.METRICS_FILE.write_text(json.dumps(self._data, indent=2), encoding="utf-8")
            self._dirty = False
            self._last_flush = now
        except Exception as exc:  # noqa: BLE001
            _host.log.debug("Metrics flush failed: %s", exc)

    # --- Recording (cheap, background-friendly) ----------------------------

    def record_tool(self, name: str, ok: bool, ms: float) -> None:
        if not _host.CFG.get("meta_learning", True) or not name:
            return
        with self._lock:
            t = self._data["tools"].setdefault(name, {"calls": 0, "ok": 0, "fail": 0, "ms_total": 0.0})
            t["calls"] += 1
            t["ok" if ok else "fail"] += 1
            t["ms_total"] += max(0.0, float(ms))
            self._flush()

    def record_response(self, path: str, ms: float) -> None:
        if not _host.CFG.get("meta_learning", True) or not path:
            return
        with self._lock:
            p = self._data["paths"].setdefault(path, {"n": 0, "ms_total": 0.0})
            p["n"] += 1
            p["ms_total"] += max(0.0, float(ms))
            self._flush()

    def record_route(self, agent: str, intent: str) -> None:
        if not _host.CFG.get("meta_learning", True):
            return
        with self._lock:
            self._last_route = {"agent": agent or "", "intent": intent or ""}
            r = self._data["routes"].setdefault(intent or "unknown", {"primus": 0, "forge": 0})
            r[agent if agent in ("primus", "forge") else "primus"] += 1
            self._flush()

    def record_feedback(self, positive: bool, note: str = "") -> str:
        """Log 👍/👎 and attribute it to the last route so routing can adapt."""
        with self._lock:
            fb = self._data["feedback"]
            fb["up" if positive else "down"] += 1
            if note:
                fb.setdefault("notes", []).append({"at": _host._now_iso(), "good": positive, "note": note[:200]})
                fb["notes"] = fb["notes"][-40:]
            # Attribute to the last route: negative on a Forge route nudges toward Primus, etc.
            intent = self._last_route.get("intent") or "unknown"
            agent = self._last_route.get("agent") or "primus"
            corr = float(self._data["route_corrections"].get(intent, 0.0))
            step = 0.25 if positive else -0.4
            # Good Forge → allow more Forge (+); bad Forge → less Forge (−); inverse for Primus.
            direction = 1.0 if agent == "forge" else -1.0
            corr += direction * step
            cap = float(_host.CFG.get("feedback_routing_max_nudge", 1.5))
            self._data["route_corrections"][intent] = max(-cap, min(cap, corr))
            self._flush(force=True)
        return "👍 logged — I'll lean into what worked." if positive else \
            "👎 logged — I'll adjust routing/approach for this kind of task."

    def route_nudge(self, intent: str) -> float:
        """Bounded score adjustment for Forge for this intent (learned from feedback)."""
        if not _host.CFG.get("adaptive_routing", True):
            return 0.0
        cap = float(_host.CFG.get("feedback_routing_max_nudge", 1.5))
        val = float(self._data.get("route_corrections", {}).get(intent or "unknown", 0.0))
        return max(-cap, min(cap, val))

    # --- Analysis / reporting ----------------------------------------------

    def weaknesses(self) -> list[str]:
        """Concrete weak spots worth improving (low tool success, slow paths, 👎 trend)."""
        out: list[str] = []
        for name, t in self._data.get("tools", {}).items():
            calls = t.get("calls", 0)
            if calls >= 4:
                rate = t.get("ok", 0) / max(1, calls)
                if rate < 0.7:
                    out.append(f"Tool `{name}` succeeds only {rate:.0%} of {calls} calls — investigate.")
                avg = t.get("ms_total", 0.0) / max(1, calls)
                if avg > 8000:
                    out.append(f"Tool `{name}` is slow (~{avg / 1000:.1f}s avg over {calls} calls).")
        for path, p in self._data.get("paths", {}).items():
            n = p.get("n", 0)
            if n >= 5:
                avg = p.get("ms_total", 0.0) / max(1, n)
                if path in ("primus", "forge") and avg > 20000:
                    out.append(f"`{path}` path averages ~{avg / 1000:.1f}s — consider tighter routing/fast-paths.")
        fb = self._data.get("feedback", {})
        if fb.get("down", 0) >= 3 and fb.get("down", 0) > fb.get("up", 0):
            out.append(f"Negative feedback trend ({fb.get('down')}👎 vs {fb.get('up')}👍) — review recent replies.")
        return out

    def report(self) -> str:
        d = self._data
        lines = ["## Performance & meta-learning"]
        tools = sorted(d.get("tools", {}).items(), key=lambda kv: kv[1].get("calls", 0), reverse=True)
        if tools:
            lines.append("\n**Top tools** (calls · success · avg):")
            for name, t in tools[:8]:
                calls = t.get("calls", 0)
                rate = t.get("ok", 0) / max(1, calls)
                avg = t.get("ms_total", 0.0) / max(1, calls)
                lines.append(f"- `{name}` — {calls} · {rate:.0%} ok · {avg / 1000:.2f}s")
        if d.get("paths"):
            lines.append("\n**Response paths** (count · avg latency):")
            for path, p in sorted(d["paths"].items(), key=lambda kv: kv[1].get("n", 0), reverse=True):
                n = p.get("n", 0)
                avg = p.get("ms_total", 0.0) / max(1, n)
                lines.append(f"- {path} — {n} · {avg / 1000:.2f}s")
        if d.get("routes"):
            lines.append("\n**Routing by intent** (Primus/Forge):")
            for intent, r in d["routes"].items():
                nudge = self.route_nudge(intent)
                tag = f" · learned nudge {nudge:+.2f}" if abs(nudge) > 0.01 else ""
                lines.append(f"- {intent}: {r.get('primus', 0)}P / {r.get('forge', 0)}F{tag}")
        fb = d.get("feedback", {})
        lines.append(f"\n**Feedback:** {fb.get('up', 0)}👍 / {fb.get('down', 0)}👎")
        try:
            lines.append(f"**Model policy:** {_host.DynamicModelManager.status()}")
        except Exception:  # noqa: BLE001
            pass
        weak = self.weaknesses()
        if weak:
            lines.append("\n**Weak spots to improve:**")
            lines += [f"- {w}" for w in weak]
        if len(lines) == 1:
            return "No usage metrics recorded yet — they accrue as we work."
        return "\n".join(lines)


def get_metrics() -> MetricsTracker:
    global _metrics_tracker
    if _metrics_tracker is None:
        _metrics_tracker = MetricsTracker()
    return _metrics_tracker


def schedule_memory_consolidation() -> None:
    global _consolidation_timer
    hours = float(_host.CFG.get("memory_consolidate_hours", 6))
    if hours <= 0:
        return

    def _run() -> None:
        try:
            msg = get_memory_system().consolidate()
            _host.log.info(msg)
        except Exception as exc:
            _host.log.warning("Memory consolidation failed: %s", exc)
        schedule_memory_consolidation()

    if _consolidation_timer:
        _consolidation_timer.cancel()
    _consolidation_timer = threading.Timer(hours * 3600, _run)
    _consolidation_timer.daemon = True
    _consolidation_timer.start()


def prepare_agent_messages(history: list) -> list[BaseMessage]:
    """Short-term memory: recent turns + optional compressed summary prefix."""
    plain = [h for h in (history or []) if isinstance(h, dict) and not h.get("metadata")]
    keep = int(_host.CFG.get("chat_recent_keep", 12))
    recent = plain[-keep:] if len(plain) > keep else plain
    msgs = _host.gradio_history_to_messages(recent)
    summary = get_memory_system()._load_chat_summary().get("summary") or ""
    if summary and len(plain) > keep:
        prefix = (
            "[Earlier in this chat (summary for continuity): "
            f"{summary[:1200]}]"
        )
        msgs = [HumanMessage(content=prefix)] + msgs
    return msgs


def retrieve_rag_context(query: str, k: Optional[int] = None) -> tuple[str, list[dict]]:
    """Backward-compatible alias — use build_memory_context for full recall."""
    return build_memory_context(query, history=None, k=k)


def build_memory_context(
    query: str,
    history: Optional[list] = None,
    k: Optional[int] = None,
    *,
    agent_id: str = "primus",
    scope: str = "",
) -> tuple[str, list[dict]]:
    """Unified proactive recall — mandatory before every Primus & Forge response.

    `scope` is an optional project-chat hint. When set, it is prepended to the RETRIEVAL
    QUERY ONLY to bias recall toward the active project (prefer, not isolate). The core
    MemorySystem recall logic is unchanged — General Chat (scope="") behaves exactly as before.
    """
    if scope:
        query = f"[{scope}]\n{query}"
    return get_memory_system().build_recall_context(
        query, history=history, k=k, agent_id=agent_id
    )


def knowledge_status_markdown() -> str:
    try:
        ms = get_memory_system().stats()
        dash = render_kb_dashboard()
    except Exception as exc:
        return f"**Knowledge base:** unavailable ({exc})"
    cons = ms.get("last_consolidation") or "never"
    return (
        dash
        + f"\n\n**Memory layer:** {ms['ltm_count']} LTM · {ms['session_count']} session · "
        f"{ms.get('archive_count', 0)} archived convos · "
        f"summary {ms['summary_chars']} chars · consolidated {cons}"
    )


def _load_interactions() -> list[dict]:
    _host.ensure_app_dirs()
    try:
        data = json.loads(_host.KB_INTERACTIONS.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError):
        return []


def _save_interactions(items: list[dict]) -> None:
    _host.KB_INTERACTIONS.write_text(json.dumps(items[-200:], indent=2), encoding="utf-8")


def _load_reflections() -> list[dict]:
    """Self-reflection log: small JSON entries on how Primus can serve the operator better."""
    _host.ensure_app_dirs()
    try:
        data = json.loads(_host.REFLECTIONS_FILE.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError):
        return []


def _save_reflections(items: list[dict]) -> None:
    cap = int(_host.CFG.get("reflection_max_entries", 80))
    _host.REFLECTIONS_FILE.write_text(json.dumps(items[-cap:], indent=2), encoding="utf-8")


def maybe_store_interaction(user_msg: str, assistant_msg: str) -> None:
    """Extract, tag, and persist memories from a completed exchange."""
    get_memory_system().process_exchange(user_msg, assistant_msg)


def start_background_index(paths: Optional[list[str]] = None) -> threading.Thread:
    """Background thread to index configured folders without blocking UI."""
    paths = paths or list(_host.CFG.get("ingest_paths") or [])

    def worker() -> None:
        _index_status["running"] = True
        _index_status["message"] = "starting…"
        try:
            kb = get_kb()
            reports = []
            for raw in paths:
                p = Path(raw).expanduser()
                if not p.exists():
                    continue
                _index_status["message"] = str(p)
                reports.append(kb.ingest_folder(p, incremental=True))
            _index_status["last"] = datetime.now().strftime("%Y-%m-%d %H:%M")
            _index_status["message"] = " · ".join(reports[:3]) if reports else "no paths"
        except Exception as exc:
            _index_status["message"] = str(exc)[:200]
            _host.log.exception("Background index failed")
        finally:
            _index_status["running"] = False

    t = threading.Thread(target=worker, daemon=True, name="primus-kb-index")
    t.start()
    return t
